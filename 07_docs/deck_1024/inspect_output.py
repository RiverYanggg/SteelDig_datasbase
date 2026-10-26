from pathlib import Path
import fitz
from PIL import Image,ImageDraw
from pptx import Presentation
B=Path(__file__).resolve().parent
pdf=B/'preview/SteelDig_数据库架构_课题三汇报_1024.pdf'
d=fitz.open(pdf)
imgs=[]
for i,p in enumerate(d):
 pix=p.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False);pix.save(str(B/f'preview/slide_{i+1:02}.png'))
 im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples);im.thumbnail((640,360));can=Image.new('RGB',(660,394),'#e8ecf2');can.paste(im,(10,24));ImageDraw.Draw(can).text((13,5),f'{i+1:02}',fill='#14385e');imgs.append(can)
for start in [0,8,16]:
 batch=imgs[start:start+8];sheet=Image.new('RGB',(1320,394*((len(batch)+1)//2)),'white')
 for k,im in enumerate(batch):sheet.paste(im,((k%2)*660,(k//2)*394))
 sheet.save(B/f'preview/montage_{start+1:02}.jpg')
p=Presentation(B.parents[1]/'SteelDig_数据库架构_课题三汇报_1024.pptx')
def normalize(t):return ''.join(t.split())
missing=[]
for i,(s,pg) in enumerate(zip(p.slides,d)):
 content=normalize(pg.get_text())
 for sh in s.shapes:
  if not sh.has_text_frame:continue
  for para in sh.text_frame.paragraphs:
   txt=normalize(para.text)
   if txt and txt not in content:missing.append((i+1,txt))
print('SLIDE PDF PAGES',len(d),'TEXT CHECK MISSING',missing)
q=fitz.open(B/'preview/SteelDig_数据库架构_逐页讲稿_1024.pdf')
print('DOCX PDF PAGES',len(q))
for i,pg in enumerate(q):
 print(i+1,len(pg.get_text()),pg.get_text().splitlines()[1:3])
