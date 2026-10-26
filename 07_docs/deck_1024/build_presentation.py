from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from io import BytesIO
import json
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from docx import Document
from docx.shared import Inches as DInches, Pt as DPt, RGBColor as DColor
from docx.oxml import OxmlElement as DOxml
from docx.oxml.ns import qn

BASE=Path(__file__).resolve().parent
ROOT=BASE.parents[1]
A=BASE/'assets'
OUT=ROOT/'SteelDig_数据库架构_课题三汇报_1024.pptx'
DOC=ROOT/'SteelDig_数据库架构_逐页讲稿_1024.docx'
FONT='Arial Unicode MS'
BLUE='225AB4'; NAVY='103B74'; INK='142B49'; MUTED='5A6D84'; LIGHT='EFF5FC'; LINE='D8E3F0'; TEAL='138B8D'; GREEN='157B63'; AMBER='AD6C15'; WHITE='FFFFFF'; PALE='F7F9FC'
prs=Presentation('/Users/mac/Desktop/课题三_1024.pptx')
body_bg=deepcopy(prs.slides[1].shapes[0]._element)
sid=prs.slides._sldIdLst[-1];prs.part.drop_rel(sid.rId);prs.slides._sldIdLst.remove(sid)
notes=[]

def rgb(c): return RGBColor.from_string(c)
def rect(s,x,y,w,h,fill=WHITE,line=None,r=.10):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
 sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
 if line: sh.line.color.rgb=rgb(line);sh.line.width=Pt(.7)
 else: sh.line.fill.background()
 if r:
  try:sh.adjustments[0]=r
  except:pass
 sh.shadow.inherit=False
 for el in sh._element.xpath('./p:style/a:effectRef'):el.set('idx','0')
 return sh

def text(s,x,y,w,h,txt,size=20,color=INK,bold=False,align=PP_ALIGN.LEFT,font=FONT,val=MSO_ANCHOR.TOP):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
 tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=tf.margin_right=Inches(.015);tf.margin_top=tf.margin_bottom=0;tf.vertical_anchor=val
 for i,line in enumerate(txt.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(5);p.line_spacing=1.15
  run=p.add_run();run.text=line;run.font.name=font;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=rgb(color)
  rp=run._r.get_or_add_rPr();ea=OxmlElement('a:ea');ea.set('typeface',font);rp.append(ea)
 return sh

def line(s,x1,y1,x2,y2,color=LINE,width=1.5,arrow=False):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=rgb(color);sh.line.width=Pt(width)
 for el in sh._element.xpath('./p:style/a:effectRef'):el.set('idx','0')
 if arrow:
  el=OxmlElement('a:tailEnd');el.set('type','triangle');sh._element.spPr.get_or_add_ln().append(el)
 return sh

def arrow(s,x,y,w=.35,h=.20,color=BLUE):
 sh=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(color);sh.line.fill.background();return sh

def pic(s,name,x,y,w,h):
 p=A/name if not Path(name).is_absolute() else Path(name)
 iw,ih=Image.open(p).size;ratio=min(w/iw,h/ih);pw,ph=iw*ratio,ih*ratio
 return s.shapes.add_picture(str(p),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph))

def card(s,x,y,w,h,title,body,accent=BLUE,size=18):
 rect(s,x,y,w,h,PALE,LINE)
 rect(s,x,y,.055,h,accent,r=0)
 text(s,x+.22,y+.20,w-.44,.48,title,21,accent,True)
 text(s,x+.22,y+.83,w-.44,h-.93,body,size,MUTED)

def chip(s,x,y,w,txt,fill=LIGHT,color=BLUE):
 rect(s,x,y,w,.36,fill,r=.16);text(s,x+.05,y+.045,w-.1,.27,txt,11,color,True,PP_ALIGN.CENTER)

def banner(s,txt,y=6.24,color=NAVY):
 rect(s,.86,y,11.6,.50,LIGHT,r=.10);text(s,1.03,y+.10,11.26,.34,txt,17,color,True)

def lead(s,txt):text(s,.88,1.38,11.55,.55,txt,24,NAVY,True)

def page(title,section,source='架构说明书（白话版） · 2026-10-23 文档快照'):
 s=prs.slides.add_slide(prs.slide_layouts[0])
 for sh in list(s.shapes):s.shapes._spTree.remove(sh._element)
 s.shapes._spTree.insert_element_before(deepcopy(body_bg),'p:extLst')
 s.shapes.add_picture(str(A/'template_2_1.png'),0,0,width=prs.slide_width,height=Inches(1.26))
 text(s,.76,.29,7.13,.65,title,26,INK,True,val=MSO_ANCHOR.MIDDLE)
 text(s,.80,7.16,9.3,.20,f'STEELDIG  /  {section}     ·     {source}',8.5,MUTED)
 text(s,11.68,7.12,.85,.28,f'{len(prs.slides):02d}',12,BLUE,True,PP_ALIGN.RIGHT)
 return s

def note(title,seconds,script,cue,source):notes.append(dict(page=len(prs.slides),title=title,seconds=seconds,script=script,cue=cue,source=source))

def table(s,x,y,width,colwidths,headers,rows,rowh=.68,size=17):
 # Native shapes and text keep every cell editable.
 xs=[x]
 for v in colwidths:xs.append(xs[-1]+width*v)
 for ri,row in enumerate([headers]+rows):
  fill=NAVY if ri==0 else (PALE if ri%2 else WHITE)
  rect(s,x,y+ri*rowh,width,rowh,fill,r=0)
  for j,val in enumerate(row):
   text(s,xs[j]+.13,y+ri*rowh+.12,width*colwidths[j]-.24,rowh-.17,str(val),size if ri else size-1,WHITE if ri==0 else INK,ri==0)
  if ri:line(s,x,y+(ri+1)*rowh,x+width,y+(ri+1)*rowh,LINE,.6)

# 01 Template cover: retain institute/BAOWU marks and original presenter.
s=prs.slides[0]
text(s,2.2,3.45,8.95,.43,'SteelDig 数据库架构与应用进展',25,WHITE,True,PP_ALIGN.CENTER)
rect(s,5.12,5.12,3.10,.46,'315995',r=.08)
text(s,5.15,5.18,3.04,.31,'2026年10月24日',21,WHITE,True,PP_ALIGN.CENTER)
note('课题三进展汇报｜SteelDig 数据库架构与应用进展',35,
'今天汇报 SteelDig 数据库建设的阶段进展。我们已经把 68 篇低密度钢相关文献的抽取结果组织成一套可以查询、核对和修订的数据底座。接下来主要讲三件事：目前形成了什么成果，为什么采用这样的架构，以及这些数据怎样支持材料研究。页面上的规模与质量数字采用 10 月 23 日的项目文档和报告快照。',
'封面沿用课题三模板；先说明主题和数字口径。','原始模板；README.md；04_reports/qa_report.md')

# 02 Snapshot
s=page('建设成果：从论文到可用数据','01 建设成果','README.md · qa_report.md（ingest_id=52）')
lead(s,'以文献事实为核心，形成可查询、可复核的数据底座')
for i,(v,u,sub) in enumerate([('68','篇文献','低密度钢相关语料'),('638','个样品','成分 × 工艺状态'),('1,508','条性能值','其中 1,304 条可数值化')]):
 x=.89+i*3.91;rect(s,x,2.18,3.72,1.85,LIGHT,r=.08);text(s,x+.22,2.40,3.25,.75,v,43,BLUE,True);text(s,x+.24,3.22,3.24,.35,u,19,NAVY,True);text(s,x+.24,3.66,3.26,.25,sub,12,MUTED)
for i,(v,l) in enumerate([('250','个合金'),('10,316','个证据块'),('68,944','条溯源记录'),('5,099','条检索文档')]):
 x=.96+i*2.92;text(s,x,4.55,2.72,.52,v,29,NAVY,True);text(s,x,5.19,2.72,.35,l,17,MUTED)
banner(s,'成果重点：数据结构化、查询入口统一、证据与人工修订有据可查')
note('建设成果：从论文到可用数据',55,'先看已经形成的成果。68 篇论文中组织出 250 个合金、638 个样品和 1,508 条性能值，其中 1,304 条可数值化。样品不是只按成分区分，同样成分经历不同工艺，也要作为不同样品记录。支撑这些事实的还有 10,316 个原文证据块、68,944 条溯源记录和 5,099 条检索文档。这里的溯源记录覆盖样品的多类字段，不能把它理解成每条性能都对应固定数量的证据。成果的价值在于把数据与证据、查询和修订连接起来。','优先讲 68 → 638 → 1,508，再补充证据规模。','README.md 数据规模；04_reports/qa_report.md')

# 03 five questions
s=page('建设目标：五项核心能力','01 建设成果')
lead(s,'从材料研究的日常问题，反推数据库需要的能力')
items=[('存得住','原始材料能否保留？','内容寻址资产湖\n结构化事实源'),('查得到','怎样找到目标材料？','成分与性能筛选\n关键词与语义检索'),('改得动','修订能否长期有效？','人工校订留痕\n入库后自动重放'),('说得清','数值来自哪份证据？','论文与证据定位\n来源批次可追踪'),('用得上','研究者怎样使用？','浏览页面与 API\n面向分析和 AI')]
for i,(t,q,b) in enumerate(items):
 x=.88+i*2.35;rect(s,x,2.24,2.20,3.53,PALE,LINE);text(s,x+.18,2.47,1.84,.38,f'0{i+1}',16,TEAL,True);text(s,x+.18,3.05,1.84,.46,t,23,BLUE,True);text(s,x+.18,3.74,1.86,.65,q,17,INK,True);line(s,x+.18,4.49,x+2.02,4.49);text(s,x+.18,4.75,1.86,.84,b,16,MUTED)
banner(s,'五项能力贯穿存储、检索、校订和应用，构成持续使用的基础')
note('建设目标：五项核心能力',45,'数据库的目标可以归纳成五项能力。存得住，是保留原始材料并稳定组织数据；查得到，是能够按成分、性能或研究问题查找；改得动，是人工修订可以保留下来；说得清，是能解释数值来自哪篇论文、哪段证据；用得上，是让材料研究者通过页面和接口使用。这五项能力既包含查询，也包含数据长期维护，后面的架构就是它们的落点。','每个词只解释一句，建立全篇路线。','00_ARCHITECTURE_PLAIN.md §0、§4')

# 04 architecture generated
s=page('总体架构：四层分工协同','02 架构设计')
pic(s,'01_architecture.png',.86,1.38,7.65,5.10)
for y,t,b in [(1.65,'消费层｜让人和程序使用','10 个页面 · 22 个接口'),(2.83,'检索层｜让事实被找到','精确 / 全文 / 模糊 / 语义'),(4.02,'事实源｜组织材料事实','PostgreSQL · 列 + JSONB + trace'),(5.19,'资产湖｜保留派生资产','本地目录 · 内容指纹 · 只增写入')]:
 text(s,8.61,y,3.68,.48,t,19,NAVY,True);text(s,8.61,y+.56,3.68,.60,b,15,MUTED)
text(s,1.12,6.57,11.12,.27,'分层价值：界面可演进，索引可重建，原始资产与校订记录需独立保存。',15,BLUE,True)
note('总体架构：四层分工协同',70,'这张图从下往上看。资产湖保留解析与抽取产生的 JSON、图片，使用内容指纹管理；事实源把它们组织为数据库表，同时保留原始 JSON；检索层为不同查询方式建立索引；消费层向研究者和程序提供页面与接口。溯源、校订、质量治理贯穿这些层。把变化速度不同的东西分开，才能独立改页面、升级检索或重新入库。还要注意：当前湖后端是本地目录，并没有部署对象存储；资产湖之外的原始 PDF、清单，以及数据库中的校订记录，也要纳入保存策略。','沿图自下而上讲；不要把图里的档案柜理解为已经具备异地备份。','00_ARCHITECTURE_PLAIN.md §2；README.md 分层一览；06_db/etl/steeldig_db/lake.py')

# 05 pipeline
s=page('数据流水线：每一步都有产物','02 架构设计')
lead(s,'解析 → 抽取 → 入库，形成可以检查和重跑的处理链')
pic(s,'02_pipeline_crop.png',.88,2.03,11.55,3.45)
for i,(t,b) in enumerate([('整篇事务','单篇要么成功，要么回滚'),('重复可重跑','相同输入避免重复累积'),('只读定稿','output/ 入库，过程日志留档')]):
 x=.94+i*3.90;text(s,x,5.64,3.70,.35,t,18,BLUE,True);text(s,x,6.10,3.70,.42,b,16,MUTED)
note('数据流水线：每一步都有产物',55,'从 PDF 到数据库，主要经过版面解析、结构化抽取和校验入库。解析产生文字、表格、图片及版面信息，抽取进一步整理成成分、工艺、性能等 JSON，入库阶段进行数据转换和关联。各阶段产物都会保留，便于从适当环节重跑。工程上有三条规则：单篇论文在同一事务中写入，避免半篇数据；重复入库避免重复累积；只把 output 下的定稿当作事实输入，过程日志单独留档。历史文档记录的约 30 秒针对当时的入库规模，不包括重新解析和模型抽取。','指出下方留存产物；无需解释工具命令。','00_ARCHITECTURE_PLAIN.md §5；ingest.py load_paper')

# 06 storage
s=page('可靠存储：资产与事实分离','02 架构设计')
lead(s,'湖保留原始派生资产，仓提供可查询的材料事实')
card(s,.9,2.17,5.12,3.28,'资产湖｜保存依据','6,667 个对象 · 148 MB\nJSON、图片按 SHA-256 寻址\n可检查内容变化，重复内容去重\n当前采用本地文件系统',BLUE,18)
card(s,7.28,2.17,5.12,3.28,'事实源｜支撑查询','36 张表 + 12 个视图 · 332 MB\n结构化列承载筛选与关联\npayload 保存原始 JSON\n由保留的输入和清单重建',TEAL,18)
arrow(s,6.28,3.36,.67,.33);text(s,6.12,3.99,1.02,.57,'重放\n入库',15,MUTED,False,PP_ALIGN.CENTER)
banner(s,'重建依赖完整输入；原始 PDF、清单和人工校订记录也需要备份')
note('可靠存储：资产与事实分离',55,'资产湖与事实源承担不同职责。湖中有 6,667 个对象、约 148 MB，主要是 JSON 和图片；事实源约 332 MB，包含 36 张表和 12 个视图。内容指纹能够检测文件内容变化，也帮助去重，但它本身不是访问控制或备份机制。事实层可以根据保留的输入与清单重新入库，湖则保存证据资产。PDF 原文仍留在 01_raw 中。完整恢复还需要考虑来源清单、校订记录和必要配置，不能只备份资产湖就认为万事具备。','强调哈希用于校验；不要说“有哈希就绝对无法篡改”。','00_ARCHITECTURE_PLAIN.md §2、§4；README.md；07_CRUD.md')

# 07 schema diagram native
s=page('数据模型：以样品连接材料信息','02 架构设计')
lead(s,'一个样品 = 一组成分 × 一套工艺状态')
# Edges first, then nodes
for a,b in [((2.5,3.4),(5.05,3.4)),((7.6,3.4),(9.95,2.65)),((7.6,3.4),(9.95,4.2)),((6.32,3.96),(6.32,5.0)),((2.5,3.96),(2.5,5.0))]:line(s,*a,*b,BLUE,2,True)
for x,y,w,h,t,b,c in [(1.15,2.75,2.7,1.2,'论文 / 合金','68 篇 / 250 个',BLUE),(4.95,2.75,2.7,1.2,'样品','638 个',TEAL),(9.25,2.09,2.75,1.16,'性能','1,508 条',BLUE),(9.25,3.70,2.75,1.16,'工艺步骤','2,610 条',BLUE),(4.95,5.0,2.7,1.16,'相与组织','1,493 条相记录',BLUE),(1.15,5.0,2.7,1.16,'原文与图表','10,316 个证据块',BLUE)]:
 rect(s,x,y,w,h,LIGHT,LINE);text(s,x+.12,y+.17,w-.24,.40,t,22,c,True,PP_ALIGN.CENTER);text(s,x+.12,y+.70,w-.24,.29,b,15,MUTED,False,PP_ALIGN.CENTER)
text(s,8.37,5.43,3.72,.60,'事实与出处通过\n论文、样品及字段路径关联',17,MUTED)
text(s,1.1,6.56,11.15,.27,'同样成分的不同热处理状态分别建样品，避免把不同实验条件的性能混在一起。',15,BLUE,True)
note('数据模型：以样品连接材料信息',60,'数据模型以样品为中心。论文提供来源，合金描述成分，而样品还包含具体的工艺状态。例如相同成分在铸态和时效态下，强度与延伸率可能明显不同，必须分别建样品。每个样品向下连接性能、工艺步骤、相和组织信息，再通过字段路径和证据块回到论文。这张图是为讲解简化的关系图，并不是所有表和外键的完整 ER 图。这样的组织方式支持跨样品分析，同时保留实验上下文。','从样品向外讲，突出工艺状态的重要性。','00_ARCHITECTURE_PLAIN.md §6；01_SCHEMA.md；qa_report.md')

# 08 design columns
s=page('字段设计：可筛选，也保留原文','02 架构设计')
lead(s,'常用条件进入结构化列，复杂信息保存在 JSONB')
card(s,.92,2.18,5.52,2.84,'结构化列｜用于筛选与关联','身份：论文号、样品号\n维度：年份、性能种类、数值、单位\n用途：比较大小、排序、跨表连接',BLUE,19)
card(s,6.66,2.18,5.74,2.84,'JSONB｜保留完整信息','payload 保存抽取原文\n容纳变化较多的条件与描述\n新增需求可从保留字段中回填',TEAL,19)
rect(s,.94,5.31,11.46,.63,NAVY);text(s,1.17,5.47,11.00,.35,'原文 “685 MPa”    →    数值 685 + 单位 MPa    ｜    原始表达同时保留',19,WHITE,True)
banner(s,'保留原文与标准化值，让精确查询和后续复核能够同时成立')
note('字段设计：可筛选，也保留原文',55,'字段设计采用结构化列与 JSONB 配合。身份标识、外键和常用筛选维度，适合做成有类型的列，比如年份、性能种类、数值和单位。不同论文差异较大的描述和条件，则保存在 payload 中，避免为每一种表达都增加新列。例如原文的 685 MPa，一方面拆成可比较的数值和标准单位，另一方面保留原始表达。以后增加查询需求时，可以从保留的信息中补字段。这是查询便利和信息保真之间的平衡。','用 685 MPa 的例子说明两种表示共存。','00_ARCHITECTURE_PLAIN.md §6.4；030_core.sql')

# 09 long table
s=page('组织方式：稳定身份与性能长表','02 架构设计')
lead(s,'一行记录一个性能值，新增性能类型无需扩展一排空列')
chip(s,.95,2.13,5.95,'示意数据｜用于解释表结构，不作为实验结果引用')
table(s,.95,2.72,7.55,[.24,.33,.20,.23],['样品','性能','数值','单位'],[['A / 铸态','屈服强度','880','MPa'],['A / 铸态','抗拉强度','1,016','MPa'],['A / 铸态','延伸率','39','%']],rowh=.67,size=17)
card(s,8.79,2.20,3.59,3.85,'身份必须带论文号','论文内：\n论文号 + 样品号\n\n跨论文：\n受控词表挂标签\n暂不自动合并实体',TEAL,18)
text(s,1.07,5.71,7.25,.43,'查询某个样品的全部性能时，再按需要组合展示。',17,MUTED)
banner(s,'事实与出处分开存储，以同一论文、样品和字段路径建立关联')
note('组织方式：稳定身份与性能长表',55,'性能采用长表，每一行表示某个样品的一种性能。页面用简化示意展示同一样品的屈服强度、抗拉强度和延伸率。真实记录还会保留误差、条件和原始表达。这样增加新的性能类型主要是增加记录，而不是不停增加列。身份也要稳定：样品编号只在论文内部有意义，因此关联时必须带论文号。跨论文先使用标准词表做分类，不把相似名称自动视为同一个材料。代价是展示一个样品的完整性能时，需要再组合查询结果。','表中去掉了误差等细节，仅解释长表；不要让听众把示意表当原始证据。','00_ARCHITECTURE_PLAIN.md §6.1、§6.3、§6.5、§6.6')

# 10 retrieval
s=page('检索设计：匹配不同的研究问法','03 查询与应用')
lead(s,'结构化条件负责筛选，文本检索负责发现相关证据')
for i,(t,q,b,c) in enumerate([('精确筛选','Al 8–11 wt%\n抗拉强度 > 800 MPa','结构化列与索引\n返回满足条件的记录',BLUE),('关键词检索','TWIP\nκ-碳化物','全文 + 模糊匹配\n找到相关文字片段',TEAL),('语义检索','“加铝为什么\n能降低密度？”','384 维向量\n按语义相近程度召回',BLUE)]):
 x=.91+i*3.89;card(s,x,2.19,3.67,3.55,t,q,c,22);text(s,x+.22,4.65,3.23,.84,b,17,MUTED)
banner(s,'文本检索以 RRF 融合全文、模糊与向量排名；精确条件单独承担约束')
note('检索设计：匹配不同的研究问法',65,'研究者有不同问法。成分区间、性能阈值等明确条件，由结构化筛选承担；关键词检索可以寻找 TWIP 或某类碳化物的相关片段；语义检索则面向自然语言研究问题，按意思相近的程度召回材料。三路文本召回是全文、模糊和向量，它们使用 RRF 合并名次信息；准确公式是各路 1 除以 60 加名次，再求和，不是简单把名次相加。语义相似度也不代表科学结论成立，最终还需要查看证据。中文查询与入库采用一致的分词处理。','精确过滤与文本召回分开解释；公式只在追问时补充。','00_ARCHITECTURE_PLAIN.md §7；050_search.sql（RRF 实现）')

# 11 composition
s=page('应用示例：按成分寻找候选材料','03 查询与应用','项目成分检索页面截图 · 2026-10-23 留档')
lead(s,'输入元素区间，先返回精确匹配，再查看近邻候选')
pic(s,'composition_crop.png',.96,2.13,7.43,4.30)
for y,num,t,b in [(2.28,'01','设定条件','选择 Fe / Mn / Al / C\n设置质量分数区间'),(3.66,'02','解释结果','精确匹配与近邻分开\n近邻说明偏离了多少'),(5.03,'03','标记推断','Fe 余量推算带标记\n引用时说明推断来源')]:
 chip(s,8.70,y,.52,num);text(s,9.40,y-.01,2.75,.39,t,21,NAVY,True);text(s,8.70,y+.53,3.52,.78,b,18,MUTED)
note('应用示例：按成分寻找候选材料',55,'成分检索是材料研究者很直接的入口。页面通过元素周期表选元素，再设质量分数区间。查询先给出精确匹配；如果没有合适结果，近邻候选会明确说明偏离了哪个元素、差了多少，而不会伪装成满足条件。Fe 经常在论文里写成余量，系统会根据其他元素推算，并做推断标记。项目文档记录 243 个 Fe 基合金中，4 条 Fe 含量为原值，239 条为推断值。使用这些数值时，要保留推断来源的说明。','这是一张已有页面截图，不把它描述为本次在线查询。','07_docs/screenshots/11-composition-slide.jpg；06_QUALITY.md §4')

# 12 sample screenshot
s=page('应用示例：在样品页核对上下文','03 查询与应用','项目样品页面截图 · 2026-10-23 留档')
lead(s,'把成分、工艺、性能和证据放在同一查看路径中')
pic(s,'sample_crop.png',.92,2.05,7.92,4.64)
for y,t,b in [(2.35,'看材料','成分和样品状态\n确认是否是目标对象'),(3.78,'看过程','工艺步骤与组织\n保留实验条件差异'),(5.18,'看依据','性能值附证据入口\n弱溯源状态直接标记')]:
 rect(s,9.08,y,.07,.92,BLUE,r=0);text(s,9.32,y-.02,2.75,.41,t,22,NAVY,True);text(s,9.32,y+.47,2.80,.70,b,17,MUTED)
note('应用示例：在样品页核对上下文',50,'进入样品页后，先确认成分和样品状态，再查看工艺、组织和性能。研究上最容易出问题的是把相似成分、不同处理状态的数据直接混用，因此这里把上下文放在同一页面。性能值旁边提供证据入口；能够定位原文的记录可以继续核对，定位不足的则标出弱溯源。页面同时保留原始数据视角，帮助使用者判断哪些值适合比较。截图展示的是已有功能，用于汇报时稳定说明操作路径。','按“材料—过程—依据”的顺序引导视线。','07_docs/screenshots/03-sample-top.jpg；00_ARCHITECTURE_PLAIN.md §4')

# 13 evidence
s=page('证据追踪：从性能值回到论文','04 数据治理','样品抽取文件及 evidence_blocks.json · DOI 10.1002/srin.202400475')
lead(s,'每条性能值可定位论文，75% 可进一步定位段落或表格')
chain=[('性能值','880 ± 19 MPa\n铸态 · 屈服强度'),('字段与证据','样品 + 字段路径\n证据块 b0035 / b0049'),('原文定位','正文说明 + 表格行\n回到对应论文')]
for i,(t,b) in enumerate(chain):
 x=.94+i*3.91;card(s,x,2.14,3.57,1.67,t,b,BLUE,17)
 if i<2:arrow(s,x+3.64,2.91,.22,.17)
rect(s,.96,4.10,7.24,1.62,PALE,LINE)
text(s,1.18,4.29,6.76,.26,'原文节选 · b0035',13,BLUE,True)
text(s,1.18,4.77,6.76,.65,'“…the yield strength from 880 MPa in the as-cast\nstate to 1041 MPa…”',19,INK,font='Arial')
text(s,8.67,4.13,3.50,.49,'1,133 条 / 75%',26,TEAL,True);text(s,8.67,4.72,3.50,.32,'具备段落或表格级证据',16,MUTED)
text(s,8.67,5.26,3.50,.45,'375 条 / 25%',24,AMBER,True);text(s,8.67,5.80,3.50,.32,'弱溯源，需回论文核对',16,MUTED)
banner(s,'68,944 条溯源覆盖多类样品字段；75% 的分母是 1,508 条性能值')
note('证据追踪：从性能值回到论文',70,'这里展示一个实际例子：铸态样品的屈服强度为 880 正负 19 MPa。抽取记录能够关联正文 b0035 和表格 b0049。正文说明铸态到时效后的变化，而表格包含 880 正负 19 的完整报告值，所以误差项也需要核对表格，不能只看正文截句。论文 DOI 为 10.1002/srin.202400475。全库 1,508 条性能值都可以定位到论文，其中 1,133 条能够进一步定位段落或表格，375 条明确标为弱溯源。75% 是证据覆盖率，既不是数值准确率，也不是全部溯源记录的覆盖率。','主动说清分母；正文引句不包含 ±19，误差来自表格。','_reframe_notes.md §5；02_extraction/10.1002_srin.202400475/preprocess/evidence_blocks.json；样品 data.json')

# 14 curation
s=page('人工校订：重跑后仍然有效','04 数据治理','07_CRUD.md · curate.py / cli.py')
pic(s,'03_curation.png',.86,1.39,7.65,5.10)
text(s,8.72,1.70,3.55,.54,'把修订单独记下来',23,NAVY,True)
for y,t,b in [(2.58,'可审计','记录目标、前后值、理由\n操作人和时间'),(3.79,'可重放','重新入库完成后\n自动应用有效校订'),(4.98,'可撤销','更新可恢复旧值\n撤销删除需重入库')]:
 text(s,8.72,y,3.54,.40,t,20,TEAL,True);text(s,8.72,y+.48,3.54,.72,b,17,MUTED)
text(s,1.11,6.57,11.16,.27,'新增论文走入库流程；增、改、删在可写白名单内进行，修订记录独立保存。',15,BLUE,True)
note('人工校订：重跑后仍然有效',65,'人工修订通过独立的校订记录管理。上方原始资产重新生成事实数据，下方校订记录再自动重放到事实层，然后页面读取当前结果。每条校订记录包含目标字段、前后值、理由、操作人和时间，避免直接改表后被重跑覆盖。新增论文通过入库流程，人工增改删则受白名单限制。撤销也有具体边界：更新可以恢复旧值；如果此前执行了删除，撤销后还需要重新入库恢复事实行。当前校订立即生效，尚未具备多人审批流程，因此多人协作治理仍是下一阶段工作。','强调底部日志独立保存；撤销删除不能承诺一键完整恢复。','07_CRUD.md §2、§4、§5、§6；curate.py')

# 15 version honest
s=page('来源管理：区分批次与历史版本','04 数据治理','paper_runs.json · 030_core.sql · ingest.py')
lead(s,'已具备来源批次追踪，完整历史保留仍需完善入库策略')
for i,(run,count) in enumerate([('0915 / pro','13 篇'),('0917 / pro','4 篇'),('0919 / pro','4 篇'),('1009 / flash','47 篇')]):
 x=.94+i*2.91;rect(s,x,2.18,2.73,1.13,LIGHT);text(s,x+.14,2.39,2.45,.33,run,19,BLUE,True,PP_ALIGN.CENTER);text(s,x+.14,2.88,2.45,.27,count,15,MUTED,False,PP_ALIGN.CENTER)
card(s,.95,3.72,5.52,2.13,'当前已实现','run_id 记录抽取来源\n校订前后值、理由与操作记录可查',BLUE,19)
card(s,6.70,3.72,5.68,2.13,'下一步补齐','现有 ETL 按论文整体重建\n多模型历史留存与对比尚需实现',TEAL,19)
banner(s,'4 个来源批次覆盖 68 篇论文，不等同于每篇论文保留了 4 个历史版本')
note('来源管理：区分批次与历史版本',55,'来源批次与历史版本需要分开理解。当前 68 篇论文来自四个抽取批次，分别覆盖 13、4、4 和 47 篇，run_id 能追踪它们的来源。表结构中也预留了按 run 和 is_latest 管理版本的能力。但是核对当前入库实现后可以看到，系统会先删除该篇论文并级联清理事实，再写入新结果，因此不能承诺同一论文多次重跑的全部历史事实已经保留。人工校订日志是另一类历史，当前可以记录和审计。后续应补充按批次保存历史的写入策略，再开展可靠的多模型对比。','本页是对原讲稿口径的核对修正，不讲成软件已经修复。','03_manifests/paper_runs.json；030_core.sql sample_data；ingest.py load_paper')

# 16 consumption
s=page('应用出口：页面、接口与研究流程','03 查询与应用','README.md · API 源码 · 检索页截图')
lead(s,'同一事实底座，支撑人工浏览、程序调用和证据检索')
pic(s,'search_crop.png',.95,2.10,7.03,4.30)
for y,t,b in [(2.27,'10 个页面','样品、论文、检索、分析与质量'),(3.57,'22 个接口','为脚本和应用提供结构化数据'),(4.87,'面向 AI 的证据入口','先检索事实与出处，再组织回答')]:
 text(s,8.43,y,3.80,.45,t,22,BLUE,True);text(s,8.43,y+.57,3.80,.74,b,17,MUTED)
text(s,1.13,6.59,11.08,.26,'训练数据专用导出流程尚未完成；证据召回能力不等于已经交付完整问答 Agent。',14,AMBER,True)
note('应用出口：页面、接口与研究流程',55,'消费层同时面向人和程序。10 个页面支持浏览、检索和分析，22 个接口提供结构化访问。截图展示自然语言检索的入口，用户可以从问题找到相关证据，再返回样品或文献核对。对于 AI 应用，这样的底座可以让回答引用检索到的事实和出处。但要区分已经完成的能力与后续产品：目前有证据检索和接口，不应直接把它说成已经交付了完整的问答 Agent；用于模型训练的专用数据导出流程也还没有完成。','网页、接口、AI 入口分别一句；勿承诺未演示的完整 Agent。','README.md；00_ARCHITECTURE_PLAIN.md §2.2、§10；05-search.jpg')

# 17 quality
s=page('质量治理：通过检查，也保留问题','04 数据治理','qa_report.md（ingest_id=52） · 06_QUALITY.md')
lead(s,'工程一致性和证据完整性，用不同指标分别衡量')
for i,(n,l,c) in enumerate([('0','error 级问题',TEAL),('844','warn 级问题',AMBER),('99.94%','证据 ID 可解析率',BLUE),('95','项测试：文档记录通过',BLUE)]):
 x=.95+i*2.91;rect(s,x,2.16,2.72,1.46,PALE,LINE);text(s,x+.15,2.38,2.42,.57,n,34,c,True);text(s,x+.15,3.10,2.42,.32,l,13,MUTED)
labels=[('图观察项缺文字证据',821),('证据 ID 未解析告警',15),('章节字段泄漏告警',8)]
for i,(t,n) in enumerate(labels):
 y=4.02+i*.58;text(s,1.01,y,3.8,.36,t,17,INK);rect(s,5.04,y+.05,5.57,.21,LIGHT,r=0);rect(s,5.04,y+.05,max(.06,5.57*n/844),.21,BLUE if i==0 else TEAL,r=0);text(s,10.88,y-.02,1.18,.35,str(n),18,BLUE,True,PP_ALIGN.RIGHT)
banner(s,'0 error 表示通过既定工程规则；它不代表文献抽取已经达到 100% 科学准确')
note('质量治理：通过检查，也保留问题',65,'质量指标要按各自口径解释。保存的报告中 error 级为零，warn 级共有 844 条，主要分为 821 条图观察项缺证据、15 条证据 ID 未解析告警和 8 条章节字段泄漏告警。证据 ID 的可解析率是 99.94%，它衡量的是标识对齐，不是性能值准确率。项目文档记录 95 项测试通过，这次 PPT 制作没有重新启动服务跑全套测试。因此页面明确写成文档记录。零工程错误说明满足既定一致性检查，但抽取内容仍然需要科学核验和持续补充。','95 项是文档快照；99.94% 与前面的 75% 分母完全不同。','04_reports/qa_report.md；07_docs/06_QUALITY.md')

# 18 quality semantics
s=page('数据使用：保留数值背后的含义','04 数据治理','qa_report.md · 06_QUALITY.md')
lead(s,'保留单位、区间与推断标记，避免“能算”掩盖“不可比”')
for i,(t,b,c) in enumerate([('数值归一','1,304 / 1,508 条可数值化\n原始值与标准化值同时保留\n强度按单位归一后比较',BLUE),('定性与区间','172 条为定性描述\n区间与误差保留原始表达\n不把趋势描述强行补成数字',TEAL),('标尺与推断','硬度查询区分 HV / HRC 等标尺\nFe 余量推算标记 is_inferred\n推断值引用时说明来源',BLUE)]):card(s,.92+i*3.89,2.22,3.66,3.49,t,b,c,17)
banner(s,'另有 174 / 638 个样品缺少结构化性能，需回到上游抽取环节补齐')
note('数据使用：保留数值背后的含义',60,'能够计算不等于可以直接比较。1,508 条性能记录中，1,304 条可数值化；172 条是定性描述，这两个类别也不是把全部数据简单二分后就能解释完。区间、误差和其他不可解析表达需要各自保留。硬度尤其要区分标尺，不能因为都是数值就放在同一张强度图里。Fe 余量推断要带标记，不能冒充原文测量值。另外，174 个样品缺少结构化性能，文档抽查指出部分属于上游漏抽。后续要改进抽取覆盖率，而不是只在数据库内填补缺失数字。','不画 1,304 与 172 相加等于总数的饼图；类别口径需保持独立。','04_reports/qa_report.md 属性量；06_QUALITY.md 已知缺陷')

# 19 selection
s=page('技术选型：控制当前系统复杂度','02 架构设计')
lead(s,'PostgreSQL + pgvector，在当前规模内统一核心能力')
rect(s,4.82,2.63,3.68,2.03,NAVY)
text(s,5.00,3.04,3.32,.51,'PostgreSQL',29,WHITE,True,PP_ALIGN.CENTER,font='Arial')
text(s,5.00,3.76,3.32,.40,'+ JSONB + pgvector',17,WHITE,False,PP_ALIGN.CENTER,font='Arial')
for x,y,t,b in [(.97,2.25,'结构化查询','关系、约束、事务'),(.97,4.40,'原文保真','JSONB 与标准列并存'),(9.03,2.25,'文本检索','全文 / 分词 / 模糊'),(9.03,4.40,'语义检索','向量表示与近邻索引')]:
 card(s,x,y,3.32,1.47,t,b,BLUE,17)
for a,b in [((4.29,2.99),(4.82,3.26)),((4.29,5.13),(4.82,4.17)),((8.50,3.26),(9.03,2.99)),((8.50,4.17),(9.03,5.13))]:line(s,*a,*b,BLUE,1.5)
banner(s,'减少组件间同步与运维成本；未来是否拆分检索，依据容量、延迟和并发实测决定')
note('技术选型：控制当前系统复杂度',55,'选 PostgreSQL 的核心原因是当前阶段可以用较少组件同时覆盖结构化查询、JSON 保真、全文和模糊检索，并通过 pgvector 支持语义检索。这样减少了不同系统之间的同步、备份和监控成本。并不是其他数据库不具备这些能力，而是项目目前规模较小，统一技术栈的收益更直接。随着检索量和并发上升，可以评估外迁检索层；决策应该看真实的容量、查询延迟、索引内存和更新压力，不能把线性估算当成已经完成的压测结论。','避免泛化贬低其他数据库；强调适配当前规模。','00_ARCHITECTURE_PLAIN.md §8、§10；050_search.sql')

# 20 position
s=page('平台定位：可引用的文献事实底座','03 查询与应用')
lead(s,'围绕文献中的材料声明，保留实验条件、出处与修订记录')
card(s,.95,2.20,5.43,3.44,'SteelDig 的建设重点','论文里的成分—工艺—组织—性能\n具体样品与证据块之间的关联\n人工校订与持续质量治理\n面向材料研究的条件检索',BLUE,19)
card(s,6.65,2.20,5.73,3.44,'与外部材料数据形成互补','未来可关联 Materials Project 等\n计算数据与文献实验数据并列\n保留数据类型、方法和来源\n跨论文 / 跨库实体对齐需验证',TEAL,19)
banner(s,'核心价值：把“某篇论文报告了什么”组织成可查询、可复核、可维护的数据')
note('平台定位：可引用的文献事实底座',50,'SteelDig 的定位是面向文献材料事实的数据底座。重点是论文中的成分、工艺、组织和性能怎样对应到具体样品，又怎样找到原文证据和后续修订。与 Materials Project 等外部材料库之间可以形成互补，但不宜把所有外部数据库都简单说成没有溯源的结果库。未来可以关联计算数据和文献实验数据，并列展示它们的方法和来源。跨库对齐需要验证，尤其不能把计算值和实验值混成同一个事实记录。','强调互补，不评价外部数据库整体能力。','00_ARCHITECTURE_PLAIN.md §1、§11；05_MATERIALS_PROJECT.md')

# 21 boundaries
s=page('当前边界：走向持续服务还需补齐','05 下一步工作','00_ARCHITECTURE_PLAIN.md §10 · 07_CRUD.md §6 · 源码核对')
lead(s,'已完成本地数据底座与应用入口，生产化能力需要继续建设')
table(s,.95,2.23,11.42,[.22,.37,.41],['方向','当前状态','下一步动作'],[
 ['数据服务','本地资产湖 + 单机部署','备份恢复演练、权限与账号'],
 ['协作校订','日志留痕，修订立即生效','审批流、冲突检测、操作治理'],
 ['历史版本','来源批次可追踪','保留历史事实、开展多模型对比'],
 ['下游应用','检索与 API 可用','训练数据导出、跨库关联'],
 ['规模扩展','5,099 条检索文档','先做容量与延迟评估，再决定拆分']
],rowh=.62,size=17)
text(s,1.08,6.37,11.13,.39,'本页区分“已交付”与“待建设”；规模推算与生产成熟度不作超前承诺。',16,AMBER,True)
note('当前边界：走向持续服务还需补齐',65,'当前阶段完成的是本地数据底座、检索和应用入口，接下来要补齐持续服务能力。首先是数据安全与恢复，包括账号权限、备份和恢复演练；其次是多人协作，包括审批和并发冲突检测；第三是完善真正保留历史事实的版本写入，再做多模型对比。应用侧还需要专用训练数据导出和跨库关联。规模扩展则通过容量和延迟评估决定，而不是先引入更多组件。原材料中的 30 秒入库和更大规模推算，都只适合作为当时环境的观察与估算，不能当作生产承诺。','按数据安全、协作、版本、应用、规模五条带过。','00_ARCHITECTURE_PLAIN.md §10；07_CRUD.md §6；ingest.py')

# 22 ending roadmap
s=page('下一步：从可用走向持续可信','05 下一步工作')
lead(s,'围绕质量、治理和应用，推进三组后续工作')
for i,(n,t,b,c) in enumerate([('01','补齐数据质量','补抽缺失性能\n完善证据与单位标记',BLUE),('02','完善治理能力','历史保留与模型对比\n审批、权限、恢复演练',TEAL),('03','打通研究应用','训练数据规范化导出\n跨论文与外部数据关联',BLUE)]):
 x=.94+i*3.88;rect(s,x,2.22,3.63,2.93,PALE,LINE);text(s,x+.23,2.44,3.18,.54,n,29,c,True);text(s,x+.23,3.21,3.18,.43,t,22,NAVY,True);text(s,x+.23,3.96,3.18,.85,b,19,MUTED)
rect(s,.96,5.62,11.41,.93,NAVY);text(s,1.15,5.88,11.03,.47,'让文献数据能查、能改、能审，成为材料研究可持续使用的底座',24,WHITE,True,PP_ALIGN.CENTER)
note('下一步：从可用走向持续可信',45,'下一步分三组推进。数据质量上补抽缺失性能、补齐证据并完善单位标记；治理能力上落实历史事实保留，再进行多模型对比，同时补齐权限、审批和恢复演练；研究应用上形成规范化训练数据导出，并逐步验证跨论文与外部数据库的关联。最终希望形成的，是一套能查、能改、能审，而且可以持续维护的材料研究数据底座。我的汇报到这里，谢谢。','收尾只强调三组动作，不临时增加未经验证的进度承诺。','README.md 下一步方向；00_ARCHITECTURE_PLAIN.md §10；源码核对')

# Save without any embedded notes (including notes inherited from source template).
prs.save(OUT)
raw=OUT.read_bytes()
with ZipFile(BytesIO(raw)) as zin, ZipFile(OUT,'w',ZIP_DEFLATED) as zout:
 for item in zin.infolist():
  name=item.filename
  if name.startswith(('ppt/notesSlides/','ppt/notesMasters/')):continue
  data=zin.read(name)
  if name.endswith('.rels') or name in ('[Content_Types].xml','ppt/presentation.xml'):
   root=etree.fromstring(data)
   for el in list(root):
    tag=etree.QName(el).localname
    if (tag=='Relationship' and el.get('Type','').endswith(('/notesSlide','/notesMaster'))) or (tag=='Override' and el.get('PartName','').startswith(('/ppt/notesSlides/','/ppt/notesMasters/'))) or tag=='notesMasterIdLst':root.remove(el)
   data=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
  zout.writestr(item,data)

(BASE/'slide_notes.json').write_text(json.dumps(notes,ensure_ascii=False,indent=2))
# Speaker notes are managed ONLY in this separate Word document.
doc=Document();sec=doc.sections[0];sec.top_margin=DInches(.75);sec.bottom_margin=DInches(.7);sec.left_margin=sec.right_margin=DInches(.83)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name];st.font.name=FONT;st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
doc.styles['Normal'].font.size=DPt(12);doc.styles['Normal'].paragraph_format.line_spacing=1.45
doc.styles['Normal'].paragraph_format.space_after=DPt(10)
for nm,sz in [('Title',29),('Heading 1',23),('Heading 2',16)]:doc.styles[nm].font.size=DPt(sz);doc.styles[nm].font.color.rgb=DColor.from_string(NAVY)
head=sec.header.paragraphs[0];head.text='STEELDIG  |  课题三进展汇报  |  独立逐页讲稿';head.style='Caption'
f=sec.footer.paragraphs[0];f.alignment=2
r=f.add_run();fld=DOxml('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
doc.add_paragraph('SteelDig 数据库架构',style='Title');doc.add_paragraph('课题三汇报 · 独立逐页讲稿',style='Subtitle')
doc.add_paragraph(f'对应 PPT：{OUT.name}\n共 {len(notes)} 页｜建议讲解约 {round(sum(x["seconds"] for x in notes)/60)} 分钟\n版本日期：2026 年 10 月 24 日')
doc.add_heading('使用说明',level=2)
for t in ['本文件是唯一的演讲者讲稿文件。PPT 内不嵌入演讲者备注。','每页提供建议时间、可直接使用的讲稿、讲解提示和核对来源。时间为建议，可按会议节奏缩短。','规模与质量数据采用项目 2026-10-23 文档及保存的 QA 报告；本次未重新运行数据库测试或容量压测。功能图采用已有项目截图。','四层架构、数据流水线与校订重放三张图使用用户提供的 gpt-image-2 调用脚本生成，提示词与原图单独留存。图示用于解释架构，统计数字由可编辑文字承载。','原讲稿中“完整多版本已保留”“RRF 是名次直接相加”等表述已按现有源码修正。']:
 doc.add_paragraph(t,style='List Bullet')
doc.add_heading('缩短至 12–15 分钟',level=2);doc.add_paragraph('优先讲 P02、P04、P05、P07、P10、P11、P13、P14、P17、P21、P22；P08、P09、P15、P18、P19、P20 可简讲。')
for n in notes:
 doc.add_page_break();doc.add_heading(f'P{n["page"]:02d}  {n["title"]}',level=1);doc.add_paragraph(f'建议时长：{n["seconds"]//60}:{n["seconds"]%60:02d}  ｜  对应 PPT 第 {n["page"]} 页',style='Subtitle')
 doc.add_heading('讲稿',level=2);doc.add_paragraph(n['script']);doc.add_heading('讲解提示',level=2);doc.add_paragraph(n['cue']);doc.add_heading('依据与口径',level=2);p=doc.add_paragraph(n['source']);p.runs[0].font.size=DPt(10)
doc.add_page_break();doc.add_heading('预设问答',level=1)
qas=[('75% 与 99.94% 有什么区别？','75% 是 1,508 条性能值中能定位到段落或表格的比例；99.94% 是来源证据 ID 经过解析对齐后的可解析率。两者分母不同，都不能代替科学准确率。'),('为什么 Fe 有星号或推断标记？','许多论文只报告合金元素，Fe 写余量。系统在适用条件下用 100 减其他元素推算 Fe，243 个 Fe 基合金中 239 条是推断值，引用时要说明。'),('发现错误后能直接改数据库吗？','通过校订入口记录目标、前后值、理由和操作人，入库后自动重放。直接修改事实表可能被重跑覆盖。更新可撤销；撤销删除一般需要重新入库。'),('四个 run 是否说明每篇论文都有四版历史？','不能。当前四个来源批次分别覆盖不同论文。虽然表结构预留版本字段，当前 ETL 会整篇删除并重建，完整历史留存和跨模型对比仍需完善。'),('这次展示是否是在线验收结果？','不是。这次以项目文档、现有截图、保存的 QA 报告和源码核对制作汇报。95 项测试通过、30 秒入库等属于既有文档记录，本次未重新做在线功能验收或压测。'),('为什么没有把流程图全部做成可编辑形状？','按本次要求使用 image2 增强三个关键机制图的视觉表达；这些图为位图。页面标题、数字、说明、关系示意图和表格仍是可编辑的 PowerPoint 对象。原图与提示词保留，可继续生成和替换。')]
for qi,(q,a) in enumerate(qas):
 if qi==3:doc.add_page_break();doc.add_heading('预设问答（续）',level=1)
 doc.add_heading(q,level=2);doc.add_paragraph(a)
doc.add_page_break();appendix_start=len(doc.paragraphs);doc.add_heading('核对记录与交付说明',level=1)
for t in ['版本：删除原稿中“旧结果全部不消失、查历史去掉 is_latest 条件”的已实现承诺。依据 ingest.py load_paper 的整篇级联删除逻辑。','检索：RRF 正确表达为按排名倒数融合，源码是 SUM(1.0 / (60 + r))。结构化精确过滤与三路文本召回分开说明。','资产：内容指纹能校验与去重，不等同于物理不可篡改，也不替代权限、备份与恢复演练。重建依赖输入、清单及独立校订记录。','质量：不把零工程 error、证据 ID 解析率或测试通过数当作科学准确率；不把 172 条定性与 1,304 条可数值化误画为覆盖全部 1,508 条的互斥分组。','外部比较：不沿用对 Excel、MySQL、MongoDB、Materials Project 等能力的绝对化否定；选型围绕本项目规模与维护成本展开。','扩展：不承诺未经压测的万人规模、毫秒级响应或线性外推性能；训练数据导出、完整 Agent、审批流与生产权限均按当前状态说明。','图像：使用 /Users/mac/Desktop/image-test/generate_image.py，模型 gpt-image-2，1536×1024，high；提示词位于 07_docs/deck_1024/prompts/，图像位于 assets/。']:
 doc.add_paragraph(t,style='List Bullet')
doc.add_heading('主要来源',level=2)
for t in ['07_docs/00_ARCHITECTURE_PLAIN.md','SteelDig_数据库架构_讲演文稿.docx','SteelDig_数据库架构_讲演版.pptx（结构与示例参考）','README.md；04_reports/qa_report.md；07_docs/06_QUALITY.md；07_docs/07_CRUD.md','06_db/etl/steeldig_db/ingest.py、curate.py；06_db/migrations/030_core.sql、050_search.sql','07_docs/screenshots/；02_extraction/10.1002_srin.202400475/；03_manifests/paper_runs.json']:
 p=doc.add_paragraph(t);p.runs[0].font.size=DPt(10)
for p in doc.paragraphs[appendix_start:]:
 p.paragraph_format.space_after=DPt(5);p.paragraph_format.line_spacing=1.2
 if p.style.name not in ('Heading 1','Heading 2'):
  for r in p.runs:r.font.size=DPt(10.5)
doc.save(DOC)
print(OUT);print(DOC);print('Slides',len(prs.slides),'Speaking seconds',sum(n['seconds'] for n in notes))
