"""生成 SteelDig 数据库设计 PPT（16:9，可编辑）。

设计语言与 Web 界面一致：深海军蓝 + 蓝色主色 + 中性灰；
全部文字/表格/图表都是原生对象，可继续编辑。
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]          # database/
SHOTS = ROOT / "07_docs" / "screenshots"
OUT = ROOT / "SteelDig_数据库设计.pptx"

W, H = 13.333, 7.5
M = 0.72

INK = RGBColor(0x0B, 0x12, 0x20)
INK2 = RGBColor(0x1D, 0x29, 0x3C)
BODY = RGBColor(0x34, 0x40, 0x54)
MUTED = RGBColor(0x66, 0x70, 0x85)
FAINT = RGBColor(0x98, 0xA2, 0xB3)
LINE = RGBColor(0xE6, 0xE8, 0xEC)
BG = RGBColor(0xF6, 0xF7, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLUE = RGBColor(0x25, 0x63, 0xEB)
BLUE_SOFT = RGBColor(0xEF, 0xF4, 0xFF)
VIOLET = RGBColor(0x7C, 0x3A, 0xED)
TEAL = RGBColor(0x0E, 0x93, 0x84)
AMBER = RGBColor(0xB5, 0x47, 0x08)
GREEN = RGBColor(0x06, 0x76, 0x47)
RED = RGBColor(0xB4, 0x23, 0x18)

FONT = "PingFang SC"
MONO = "Menlo"


def _font(run, size, bold=False, color=BODY, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def textbox(slide, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return tf


def para(tf, text, size=13, bold=False, color=BODY, space_before=0, space_after=4,
         line=1.25, align=PP_ALIGN.LEFT, first=False, name=FONT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    r = p.add_run()
    r.text = text
    _font(r, size, bold, color, name)
    return p


def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
         line_w=0.75, shadow=False):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    # 去掉主题样式引用（否则 LibreOffice/PowerPoint 仍会套用主题阴影）
    for st in shp._element.findall(qn("p:style")):
        shp._element.remove(st)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        adj = radius if radius is not None else 0.06
        shp.adjustments[0] = adj
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.14)
    tf.margin_top = tf.margin_bottom = Inches(0.08)
    return shp


def new_slide(prs, bg=WHITE):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fill = s.background.fill
    fill.solid()
    fill.fore_color.rgb = bg
    return s


def head(slide, title, kicker=None, dark=False):
    y = 0.46
    if kicker:
        tf = textbox(slide, M, y - 0.03, 10, 0.3)
        para(tf, kicker.upper(), 9.5, True, BLUE if not dark else RGBColor(0x7D, 0xA7, 0xFF),
             first=True, space_after=0)
        y += 0.30
    tf = textbox(slide, M, y, W - 2 * M, 0.6)
    para(tf, title, 27, True, WHITE if dark else INK, first=True, space_after=0, line=1.1)


def footer(slide, n, dark=False, note="SteelDig · 材料知识库设计"):
    c = FAINT if not dark else RGBColor(0x5C, 0x67, 0x79)
    tf = textbox(slide, M, H - 0.46, 8, 0.25)
    para(tf, note, 9, False, c, first=True)
    tf = textbox(slide, W - M - 0.6, H - 0.46, 0.6, 0.25, PP_ALIGN.RIGHT)
    para(tf, f"{n:02d}", 9, False, c, first=True)


def rule(slide, y=1.30, color=BLUE, w=0.52):
    rect(slide, M, y, w, 0.045, fill=color, shape=MSO_SHAPE.RECTANGLE)


def kpi(slide, x, y, w, label, value, unit="", foot=None, color=INK):
    tf = textbox(slide, x, y, w, 0.26)
    para(tf, label, 10.5, False, MUTED, first=True)
    tf = textbox(slide, x, y + 0.30, w, 0.52)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = value; _font(r, 30, True, color)
    if unit:
        r2 = p.add_run(); r2.text = " " + unit; _font(r2, 12, False, MUTED)
    if foot:
        tf = textbox(slide, x, y + 0.90, w, 0.28)
        para(tf, foot, 9.5, False, FAINT, first=True)


def bullet(tf, text, first=False, size=13, color=BODY, bold=False, dot=True, gap=7):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(gap)
    p.line_spacing = 1.3
    r = p.add_run()
    r.text = ("· " if dot else "") + text
    _font(r, size, bold, color)
    p.level = 0


def code_block(slide, x, y, w, h, lines, size=10.5):
    rect(slide, x, y, w, h, fill=RGBColor(0x11, 0x18, 0x28), line=None, radius=0.03)
    tf = textbox(slide, x + 0.18, y + 0.16, w - 0.36, h - 0.32)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(1)
        p.line_spacing = 1.22
        r = p.add_run(); r.text = ln
        _font(r, size, False, RGBColor(0xD5, 0xDB, 0xE7), MONO)


def table(slide, x, y, w, rows, col_w, header=True, size=11, row_h=0.34, header_h=None,
          header_bg=RGBColor(0xF2, 0xF4, 0xF7), zebra=True, align=None):
    n_rows, n_cols = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y),
                                   Inches(w), Inches(row_h * n_rows))
    tbl = shape.table
    tbl.first_row = header
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(header_h if (i == 0 and header and header_h) else row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = Inches(0.10)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if i == 0 and header:
                cell.fill.fore_color.rgb = header_bg
            elif zebra and i % 2 == 0:
                cell.fill.fore_color.rgb = RGBColor(0xFB, 0xFC, 0xFD)
            else:
                cell.fill.fore_color.rgb = WHITE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            if align and align[j] == "r":
                p.alignment = PP_ALIGN.RIGHT
            elif align and align[j] == "c":
                p.alignment = PP_ALIGN.CENTER
            r = p.add_run(); r.text = str(val)
            if i == 0 and header:
                _font(r, size - 0.5, True, MUTED)
            else:
                _font(r, size, False, BODY)
    return tbl


def pill(slide, x, y, w, text, bg=BLUE_SOFT, fg=BLUE, size=10):
    shp = rect(slide, x, y, w, 0.28, fill=bg, line=None, radius=0.5)
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    _font(r, size, True, fg)
    return shp


def picture(slide, path, x, y, w, h=None, border=True):
    if h is None:
        pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    else:
        pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                       width=Inches(w), height=Inches(h))
    if border:
        pic.line.color.rgb = LINE
        pic.line.width = Pt(0.75)
    return pic


def arrow_down(slide, x, y, h=0.26, color=RGBColor(0xC7, 0xD2, 0xFE), w=0.18):
    shp = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x), Inches(y),
                                 Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    for st in shp._element.findall(qn("p:style")):
        shp._element.remove(st)
    return shp
