"""生成讲演文稿 DOCX，并把同一份讲稿写进 PPT 的演讲者备注。"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from pptx import Presentation

from script_content import QA, SLIDES

ROOT = Path(__file__).resolve().parents[2]
DOCX = ROOT / "SteelDig_数据库架构_讲演文稿.docx"
PPTX = ROOT / "SteelDig_数据库架构_讲演版.pptx"

INK = RGBColor(0x1D, 0x29, 0x3C)
MUTED = RGBColor(0x66, 0x70, 0x85)
FAINT = RGBColor(0x98, 0xA2, 0xB3)
BLUE = RGBColor(0x25, 0x63, 0xEB)
FONT = "PingFang SC"


def set_font(run, size=None, bold=None, color=None, name=FONT):
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


def para(doc, text="", size=11.5, bold=False, color=INK, space_after=6,
         space_before=0, line=1.55, indent=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = line
    if indent is not None:
        pf.left_indent = Cm(indent)
    r = p.add_run(text)
    set_font(r, size, bold, color)
    return p


def rule(doc, color="E6E8EC", space_before=6, space_after=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    pPr.append(pbdr)


def add_footer(section, text):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, 8, False, FAINT)


doc = Document()
sec = doc.sections[0]
sec.page_height = Cm(29.7)
sec.page_width = Cm(21.0)
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.4)
sec.right_margin = Cm(2.4)
add_footer(sec, "SteelDig · 架构与设计逻辑 · 讲演文稿")

style = doc.styles["Normal"]
style.font.name = FONT
style.font.size = Pt(11.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

# ---------------------------------------------------------------------------
# 封面
# ---------------------------------------------------------------------------
para(doc, "STEELDIG · MATERIALS KNOWLEDGE BASE", 9.5, True, BLUE, space_before=60, space_after=8)
para(doc, "讲演文稿", 34, True, INK, space_after=8)
para(doc, "把论文变成数据库：一个能查、能改、能审的材料知识库", 15, False, MUTED, space_after=24)
rule(doc, space_before=2, space_after=14)
para(doc, "25 页 · 建议 25 分钟 · 面向非技术听众", 10.5, False, MUTED, space_after=6)
para(doc, "每页三段：讲稿、时间、提示。讲稿可以直接念，也可以只记要点。", 10.5, False, MUTED, space_after=6)
para(doc, "PPT 的演讲者备注里是同一份文字，放映时按“备注”视图即可看到。", 10.5, False, MUTED, space_after=30)
para(doc, "使用建议", 10, True, INK, space_after=4)
para(doc, "· 第一次讲：通读一遍，标出自己顺口的句子，其余可以删。", 10, False, MUTED, space_after=2)
para(doc, "· 时间紧：优先保 P02、P05、P09、P12、P13、P19、P24。", 10, False, MUTED, space_after=2)
para(doc, "· 被问到答不上来的：先承认，再说“会后查了发您”，不要现场编。", 10, False, MUTED, space_after=2)
para(doc, "· 最后一页之后有 6 个预设问答，建议提前看一遍。", 10, False, MUTED, space_after=2)

# ---------------------------------------------------------------------------
# 逐页
# ---------------------------------------------------------------------------
for item in SLIDES:
    doc.add_page_break()
    para(doc, f"P{item['n']:02d}", 9.5, True, BLUE, space_after=2)
    para(doc, item["title"], 17, True, INK, space_after=3)
    para(doc, f"建议时长 {item['time']} · 对应 PPT 第 {item['n']} 页", 9, False, FAINT, space_after=4)
    rule(doc, space_before=2, space_after=10)
    for line in item["script"]:
        para(doc, line, 12, False, INK, space_after=7, line=1.75)
    para(doc, "提示", 9, True, MUTED, space_before=6, space_after=3)
    for tip in item["tips"]:
        para(doc, "· " + tip, 9.5, False, MUTED, space_after=2, line=1.5)

# ---------------------------------------------------------------------------
# 问答
# ---------------------------------------------------------------------------
doc.add_page_break()
para(doc, "附：预设问答", 20, True, INK, space_after=4)
para(doc, "这 6 个问题最可能被问到，答案都控制在一分钟内。", 10, False, MUTED, space_after=10)
for i, (q, a) in enumerate(QA, start=1):
    para(doc, f"Q{i}. {q}", 12, True, INK, space_before=8, space_after=4)
    para(doc, a, 11, False, MUTED, space_after=6, line=1.6)

doc.save(DOCX)
print(f"saved {DOCX}")

# ---------------------------------------------------------------------------
# 写进 PPT 备注
# ---------------------------------------------------------------------------
prs = Presentation(PPTX)
assert len(prs.slides) == len(SLIDES), (len(prs.slides), len(SLIDES))
for slide, item in zip(prs.slides, SLIDES):
    tf = slide.notes_slide.notes_text_frame
    tf.text = f"【讲稿 · 建议 {item['time']}】\n" + "\n".join(item["script"]) + \
              "\n\n提示：" + "；".join(item["tips"])
prs.save(PPTX)
print(f"notes written: {PPTX}")
