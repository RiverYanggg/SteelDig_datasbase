"""将当前 Markdown 分镜渲染成可交付的横向 DOCX。

用法：database/06_db/.venv/bin/python make_script_doc.py
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
MD = ROOT / "脚本分镜.md"
DOCX = ROOT.parents[1] / "SteelDig_演示视频_分镜脚本.docx"

FONT_CN = "微软雅黑"
FONT_BODY = "宋体"
FONT_MONO = "Consolas"
INK = RGBColor(0x10, 0x18, 0x28)
MUTED = RGBColor(0x66, 0x70, 0x85)
BLUE = RGBColor(0x25, 0x63, 0xEB)


def clean(text: str) -> str:
    text = re.sub(r"\[(.*?)\]\(.*?\)", r"\1", text)
    return text.replace("**", "").replace("`", "")


def shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def set_font(run, name=FONT_BODY, size=9.5, bold=False, color=INK) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), name)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)


def add_para(doc, text, size=9.5, bold=False, color=INK, style=None, mono=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(clean(text))
    set_font(r, FONT_MONO if mono else FONT_BODY, size, bold, color)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(clean(text.lstrip("# ")))
    set_font(r, FONT_CN, 16 if level == 1 else 12, True, INK if level == 1 else BLUE)
    if level == 1:
        ppr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:color"), "2563EB")
        pbdr.append(bottom)
        ppr.append(pbdr)


def add_code(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(.25)
    p.paragraph_format.space_after = Pt(6)
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F5F7FA")
    ppr.append(shd)
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        set_font(r, FONT_MONO, 8.2, False, RGBColor(0x1D, 0x29, 0x3C))


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [3.9, 7, 15.9] if len(rows[0]) == 3 else [4.5, 22.3]
    if len(rows[0]) not in (2, 3):
        widths = [26.8 / len(rows[0])] * len(rows[0])
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for j, value in enumerate(rows[0]):
        c = table.rows[0].cells[j]
        c.width = Cm(widths[j])
        shade(c, "EFF4FF")
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(clean(value))
        set_font(r, FONT_CN, 8.5, True, BLUE)
    for values in rows[1:]:
        cells = table.add_row().cells
        no_split = OxmlElement("w:cantSplit")
        table.rows[-1]._tr.get_or_add_trPr().append(no_split)
        for j, value in enumerate(values):
            cells[j].width = Cm(widths[j])
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(clean(value))
            set_font(r, FONT_BODY, 9, False, INK)
    return table


def render():
    lines = MD.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21)
    section.left_margin = section.right_margin = Cm(1.45)
    section.top_margin = section.bottom_margin = Cm(1.25)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.add_run("SteelDig · 演示分镜 v2.0    |    "), FONT_CN, 8, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    i = 0
    in_code = False
    code_lines = []
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            if in_code:
                add_code(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            add_heading(doc, line, 1 if level <= 1 else 2)
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[-: ]+", c) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                add_table(doc, rows)
            continue
        if line.startswith("- "):
            add_para(doc, line[2:], size=9, style="List Bullet")
        elif line.startswith("> "):
            add_para(doc, line[2:], size=9, color=MUTED)
        else:
            add_para(doc, line, size=9.5)
        i += 1

    doc.save(DOCX)
    print(f"saved {DOCX}")


if __name__ == "__main__":
    render()
