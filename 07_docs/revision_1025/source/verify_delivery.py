from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from pptx import Presentation
import json,hashlib,pymupdf
r=Path('/Users/mac/Desktop/SteelDig/database');w=r/'07_docs/revision_1025';p=r/'SteelDig_SFT与数据库_优化版_1025.pptx';orig=r/'SteelDig_数据库架构_课题三汇报_1024.pptx'
prs=Presentation(p);source=Presentation(orig);assert len(prs.slides)==41
with ZipFile(p) as z:
 names=z.namelist();notes=[n for n in names if n.startswith('ppt/notesSlides/notesSlide')];assert not notes,notes
 fonts=set()
 for n in names:
  if n.startswith('ppt/slides/slide') and n.endswith('.xml'):
   root=etree.fromstring(z.read(n));fonts.update(el.get('typeface') for el in root.iter() if etree.QName(el).localname in ['latin','ea','cs'])
 assert fonts=={'Microsoft YaHei'},fonts
mapping=json.loads((w/'source/slide_mapping.json').read_text())
def pictures(slide):
 return [hashlib.sha256(sh.image.blob).hexdigest() for sh in slide.shapes if sh.shape_type==13]
for m in mapping:
 if isinstance(m['original'],int):assert pictures(source.slides[m['original']-1])==pictures(prs.slides[m['page']-1]),m
assert hashlib.sha256(orig.read_bytes()).hexdigest()=='53c164035b4104f870b30570e83011b6a5825085f01910b2f030ed53c2e1cfda'
sg=json.loads((w/'source/actual_subgraph.json').read_text());qa=json.loads((w/'source/qa_case.json').read_text());n={x['id']:x for x in sg['nodes']};paths=[f"{n[e['from']]['label']} --({e['relation']})--> {n[e['to']]['label']}" for e in sg['edges']];assert paths==qa['relation_path']
mp=json.loads((w/'source/mp7029_snapshot.json').read_text());assert mp['band_gap']==5.51 and mp['poly_total']==4.04 and mp['space_group']==96 and mp['nsites']==12
notes=json.loads((w/'source/slide_notes.json').read_text());assert len(notes)==41
pdf=pymupdf.open(r/'SteelDig_SFT与数据库_优化版_1025.pdf');assert len(pdf)==41
badfonts=set()
for page in pdf:
 for f in page.get_fonts():
  if any(t in f[3].lower() for t in ['baoli','kaiti','xingkai']):badfonts.add(f[3])
assert not badfonts,badfonts
out={'slides':41,'embedded_notes':0,'editable_text_font':list(fonts),'all_original_slide_images_preserved':True,'original_file_unchanged':True,'graph_matches_retained_qa':True,'mp_snapshot_values_verified':True,'manuscript_sections':len(notes),'pdf_pages':len(pdf)}
(w/'source/validation.json').write_text(json.dumps(out,ensure_ascii=False,indent=2));print(json.dumps(out,ensure_ascii=False,indent=2))
for name,needle in [('honeybee','Table 2: Evaluation'),('melt_paper','Table 1: Evaluation'),('llamat_paper','Table E.1: F1-score')]:
 d=pymupdf.open(w/'research'/f'{name}.pdf');print(name,[i+1 for i,pg in enumerate(d) if needle in pg.get_text()])
