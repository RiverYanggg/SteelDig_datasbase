from pathlib import Path
from copy import deepcopy
import json, zipfile, hashlib
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from docx import Document
from docx.shared import Inches as DI, Pt as DP, RGBColor as DC
from docx.oxml.ns import qn

ROOT=Path('/Users/mac/Desktop/SteelDig/database')
WORK=ROOT/'07_docs/revision_1025'
ORIG=ROOT/'SteelDig_数据库架构_课题三汇报_1024.pptx'
OUT=ROOT/'SteelDig_SFT与数据库_优化版_1025.pptx'
FONT='Microsoft YaHei'
NAVY='153F83';BLUE='2468C8';TEAL='148F91';GOLD='B78028';INK='172B43';MUTED='60718A';LIGHT='F1F5FB';LINE='D8E3EF';WHITE='FFFFFF';PALETEAL='EFF8F7'
prs=Presentation(ORIG);original=list(prs.slides);added={};newnotes={};checks=[]
original_hash=hashlib.sha256(ORIG.read_bytes()).hexdigest()
sftnotes=json.loads(Path('/Users/mac/Desktop/SFT框架汇报_1024/source/slide_content.json').read_text())['notes']
dbnotes=json.loads((ROOT/'07_docs/deck_1024/slide_notes.json').read_text())

def text(s,x,y,w,h,txt,size=18,c=INK,bold=False,align=PP_ALIGN.LEFT):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_left=tf.margin_right=Inches(.02);tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(txt.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align;p.space_after=Pt(4);p.line_spacing=1.12
  rr=p.add_run();rr.text=line;rr.font.name=FONT;rr.font.size=Pt(size);rr.font.bold=bold;rr.font.color.rgb=RGBColor.from_string(c)
  ea=OxmlElement('a:ea');ea.set('typeface',FONT);rr._r.get_or_add_rPr().append(ea)
 checks.append(dict(slide_id=s.slide_id,text=txt,box=[x,y,w,h],size=size))
 return sh

def rect(s,x,y,w,h,fill=LIGHT,stroke=None):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h));sh.adjustments[0]=.10
 sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill)
 if stroke:sh.line.color.rgb=RGBColor.from_string(stroke);sh.line.width=Pt(.7)
 else:sh.line.fill.background()
 sh.shadow.inherit=False
 for el in sh._element.xpath('./p:style/a:effectRef'):el.set('idx','0')
 return sh

def arrow(s,x1,y1,x2,y2,c=BLUE,width=1.6):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=RGBColor.from_string(c);sh.line.width=Pt(width)
 el=OxmlElement('a:tailEnd');el.set('type','triangle');sh.line._get_or_add_ln().append(el)
 for e in sh._element.xpath('./p:style/a:effectRef'):e.set('idx','0')
 return sh

def pill(s,x,y,w,label,c=BLUE):
 rect(s,x,y,w,.32,LIGHT);text(s,x+.04,y+.04,w-.08,.25,label,11,c,True,PP_ALIGN.CENTER)

def banner(s,txt,y=6.13,size=15,fill=PALETEAL):
 rect(s,1.02,y,11.25,.43,fill);text(s,1.14,y+.075,11.01,.30,txt,size,TEAL,True)

def page(key,title,sub,source):
 s=prs.slides.add_slide(prs.slide_layouts[6])
 for sh in list(s.shapes):s.shapes._spTree.remove(sh._element)
 # Retain the user's rounded body panel and institutional header.
 for sh in list(original[1].shapes)[:2]:
  el=deepcopy(sh._element)
  for e in el.iter():
   for k,v in list(e.attrib.items()):
    if k.endswith('}embed'):
     rel=original[1].part.rels[v];e.set(k,s.part.relate_to(rel.target_part,rel.reltype))
  s.shapes._spTree.insert_element_before(el,'p:extLst')
 text(s,.75,.30,7.2,.56,title,26,NAVY,True)
 text(s,.98,1.40,11.35,.56,sub,17,MUTED)
 text(s,.99,6.77,10.95,.20,source,8,MUTED)
 text(s,11.95,6.74,.4,.25,'00',10,MUTED,False,PP_ALIGN.RIGHT)
 added[key]=s
 return s

def note(key,title,script,sources):newnotes[key]=dict(title=title,script=script,sources=sources)

# Two background pages immediately after SFT section cover.
s=page('background','微调背景：让模型学会材料任务','微调：在预训练模型基础上继续训练；SFT 用“问题—参考答案”更新参数或适配器。','参考：Lewis et al., RAG, NeurIPS 2020；Luo et al., 2023（灾难性遗忘）；本项目训练数据定位')
for x,title,lines,c in [
 (1.04,'SFT｜把任务能力学进模型','训练时：用 QA 示例更新参数\n适合：术语、回答格式、机制解释\n更新知识：通常需要再次训练',BLUE),
 (6.82,'RAG｜回答时检索外部证据','推理时：检索资料并加入上下文\n适合：最新事实、具体数值、出处\n更新知识：更新资料库与索引',TEAL)]:
 rect(s,x,2.08,5.45,1.83,LIGHT);text(s,x+.20,2.25,5.03,.38,title,21,c,True);text(s,x+.20,2.78,5.03,.99,lines,16,INK)
text(s,1.07,4.04,10.8,.34,'主要难点：训练得更专业，也要保持可信与通用能力',20,NAVY,True)
for x,t,b,c in [(1.05,'幻觉与错误固化','错误标签可能被学进参数；\n高质量数据、证据核验与拒答评测。',BLUE),(4.88,'灾难性遗忘','领域训练可能削弱原有能力；\n混合通用数据、控制训练并回归评测。',TEAL),(8.71,'数据稀缺与泛化','材料条件多、稀有机制少；\n按论文划分，测试未见工艺与材料。',GOLD)]:
 rect(s,x,4.57,3.56,1.24,LIGHT);text(s,x+.16,4.72,3.24,.32,t,18,c,True);text(s,x+.16,5.17,3.24,.53,b,13.5,MUTED)
banner(s,'两者可组合：SFT 训练“怎样回答”，RAG 提供“本次依据”；都需要验证事实。',size=15)
note('background','微调是什么，为什么还需要 RAG','微调是在预训练模型上继续训练。我们这里讨论的是监督微调，也就是用问题和参考答案教模型完成材料任务，可以更新全部参数，也可以只训练适配器。SFT 主要改变模型行为和任务能力；RAG 在回答时查外部资料，把证据放进上下文，通常不更新生成模型的参数。两者可以组合：模型学会按成分、工艺、组织和性能解释问题，具体数值和出处再从数据库检索。主要难点有三类。第一，微调并不会自动消除幻觉，错误答案还可能被固化。第二，集中学习领域数据可能损害通用能力，这就是灾难性遗忘，参数高效微调也不保证完全避免。第三，材料数据条件复杂且稀疏，需要按论文划分数据，并在未见材料或工艺上测试。RAG 也会检索错误或误读证据，因此二者都需要事实评测。',['https://arxiv.org/abs/2005.11401','https://arxiv.org/abs/2308.08747'])

s=page('research','材料领域研究：收益来自数据与训练','代表性工作（2023–2025）：区分指令微调、持续预训练，以及各自的评价口径。','[1] HoneyBee, arXiv:2310.08511v1, Tab.2；[2] LLaMat, arXiv:2412.09560v2, Tab.E.1；[3] MELT, EMNLP 2024, Tab.1')
rows=[
 ('HoneyBee','2023｜MatSci-Instruct 指令微调','生成、独立验证，\n再按反馈改进数据','7B：85.42 → 91.95','GPT-4 评定的 Accuracy；\n无验证版本 → Stage 3。',BLUE),
 ('LLaMat','2025 v2｜持续预训练＋指令微调','材料语料适配后进行任务训练','86.23 → 93.03','MatNLP 平均 Micro-F1（%）；\nLLaMA-3-chat FT → LLaMat-3-chat。',TEAL),
 ('MELT','2024｜持续预训练','语义图组织知识，\n逐步学习专业概念','0.693 → 0.741','MatSci-NLP 总体 Micro-F1；\nSciBERT → MELT。',GOLD)]
for i,(name,typ,method,metric,limit,c) in enumerate(rows):
 y=2.05+i*1.25;rect(s,1.04,y,11.22,1.12,LIGHT)
 text(s,1.23,y+.13,3.46,.33,name,19,c,True);text(s,1.23,y+.62,3.35,.29,typ,12,MUTED)
 text(s,4.73,y+.16,3.21,.72,method,17,INK)
 text(s,8.33,y+.10,3.64,.35,metric,22,c,True);text(s,8.33,y+.57,3.62,.49,limit,12,MUTED)
banner(s,'对本项目的启示：机制化组织知识＋审查数据；训练收益仍须用独立测试集验证。',y=6.02,size=15)
text(s,1.12,6.54,11,.17,'不同任务、底座和评价方法之间不可横向排名；这些分数不代表开放式材料问答已无幻觉。',10,MUTED)
note('research','材料领域已有研究及效果','这里选择三个代表性工作，不把不同指标排成一个排行榜。HoneyBee 的 MatSci-Instruct 使用生成器与独立验证器，并通过反馈逐轮改进指令数据。论文表 2 中，7B 无验证版本的 Accuracy 为 85.42，Stage 3 为 91.95。这是 GPT-4 对其生成测试集的评价分数，不是人工或实验认证的正确率，而且变化同时包含验证和迭代训练的作用。LLaMat 在材料语料上持续预训练并进行指令微调。2025 年 v2 的表 E.1 中，LLaMA-3-chat FT 与 LLaMat-3-chat 的 MatNLP 平均 Micro-F1 分别是 86.2250 和 93.0270，展示了领域适配的收益。MELT 不是 QA 指令微调，而是用语义图与课程式策略进行持续预训练；其表 1 中 SciBERT 与 MELT 的总体 Micro-F1 是 0.693 和 0.741。三项研究共同提示，数据组织与训练设计很重要，但不能据此推断本项目已经提升了模型能力。我们目前汇报的是数据建设，仍需独立训练与测试。',['https://arxiv.org/abs/2310.08511v1','https://arxiv.org/abs/2412.09560v2','https://aclanthology.org/2024.findings-emnlp.627/'])

# Actual graph, never invent a node to stand in for the original graph.
s=page('graph','真实案例：这条 QA 对应什么图谱','同一论文、同一样品：Fe–30.34Mn–8.93Al–0.9C；热轧后水淬，未再固溶。','依据：历史 graph/10.1016_j.corsci.2025.113331.json（path-5、r5、r9）；对应 records.jsonl')
pill(s,1.07,2.00,2.1,'选中机制路径：path-5')
text(s,3.38,2.02,8.55,.32,'节点保留原始标签；中文解释说明节点在资产中的实际含义。',14,MUTED)
nodeinfo=[(1.08,'κ-carbide','microstructure｜组织节点','纳米 κ-碳化物\n形貌：5–10 nm 析出物',BLUE),
 (5.05,'interface_S0','interface｜界面记录','γ/κ 界面、缺陷相互作用\n保留未固溶样品的相演化信息',TEAL),
 (9.02,'perf_S0_baseline','performance｜性能记录','3.5 wt% NaCl 环境\n耐蚀性比较的基线记录',GOLD)]
for x,t,typ,body,c in nodeinfo:
 rect(s,x,2.66,3.18,1.67,LIGHT,LINE);text(s,x+.12,2.82,2.94,.35,t,20,c,True,PP_ALIGN.CENTER);text(s,x+.12,3.28,2.94,.26,typ,12,c,False,PP_ALIGN.CENTER);text(s,x+.14,3.65,2.90,.58,body,12.5,INK,False,PP_ALIGN.CENTER)
arrow(s,4.29,3.39,5.00,3.39,BLUE);arrow(s,8.27,3.39,8.97,3.39,TEAL)
text(s,4.24,2.70,.85,.55,'r5\nsuppresses',9,BLUE,True,PP_ALIGN.CENTER)
text(s,8.25,2.70,.78,.55,'r9\ncauses',10,TEAL,True,PP_ALIGN.CENTER)
for x,t,b,c in [(1.08,'r5｜抑制：κ-碳化物 → 界面记录','边文本：碳化物消耗 Al，降低膜表面 Al 含量，\n使连续、致密的富 Al 保护膜难以形成。',BLUE),(6.88,'r9｜导致：界面记录 → 性能记录','边文本：形成的膜较薄，Fe / Mn 含量较高，\n保护能力较差，导致耐蚀性降低。',TEAL)]:
 rect(s,x,4.52,5.32,1.25,LIGHT);text(s,x+.16,4.67,5.0,.32,t,16,c,True);text(s,x+.16,5.10,5.0,.59,b,13.5,INK)
banner(s,'路径 = r5 ＋ r9；“富 Al 膜”在边的机制文本中，当前并不是单独的图节点。',y=6.04,size=15)
text(s,1.10,6.53,11.1,.18,'证据索引来自节点来源合并；关联可追溯不等于关系已经通过原文逐条核验。',10,MUTED)
note('graph','把真实子图展开看','前一页讲的是便于理解的科学机制，这一页展示实际保存的图谱。该 QA 对应 path-5，连接三类节点：κ-carbide、interface_S0 和 perf_S0_baseline，边是 r5 的 suppresses 和 r9 的 causes。三个节点都属于同一个热轧水淬样品。需要特别澄清，interface_S0 在原始资产中保存的是 γ/κ 界面、缺陷相互作用和相演化信息，并不是一个专门的氧化膜节点；关于富 Al 膜形成受限的机制保存在 r5 的文本中。r9 再描述膜保护能力不足与耐蚀性较差的关系。路径把这两条边文本拼接，提供出题焦点。这样既能看见实际图结构，也能看出当前节点粒度还不够细。原始节点 ID、属性、关系与证据引用均另存到配套 source 文件，便于复核。',['sft_mechanism_qa/out_sft/graph/10.1016_j.corsci.2025.113331.json：path-5 / r5 / r9','sft_mechanism_qa/out_sft/records.jsonl：mqa_mechanism_forward_10.1016_j.corsci.2025.113331_1'])

s=page('spec','从图谱到 QA：先定焦点，再写问答','图谱负责确定“解释哪段关系”；事实卡补充对象、条件和机制细节。','代码依据：qa/spec.py:build_specs → qa/generate.py:generate_variants；历史 QA 原始记录')
for x,n,t,b,c in [(1.05,'01','路径 → 出题规格','题型：mechanism_forward\n锚点：r5、r9\n焦点：Al 受限 → 成膜与耐蚀性',BLUE),(4.91,'02','样品 → 事实卡','成分：Fe–Mn–Al–C\n工艺：热轧、水淬、未再固溶\n补充：结构化属性＋Wiki 机制',TEAL),(8.77,'03','规格＋事实卡 → 生成','输入：题型＋焦点＋上下文\n要求：解释原因，交代研究对象\n约束：不输出 S0 / r5 等代号',GOLD)]:
 rect(s,x,2.10,3.49,2.05,LIGHT);pill(s,x+.15,2.25,.48,n,c);text(s,x+.15,2.78,3.19,.36,t,19,c,True);text(s,x+.15,3.25,3.19,.83,b,12.5,INK)
arrow(s,4.58,3.03,4.86,3.03);arrow(s,8.44,3.03,8.72,3.03,TEAL)
rect(s,1.05,4.30,11.21,.90,PALETEAL)
text(s,1.20,4.45,.45,.35,'Q',22,TEAL,True)
text(s,1.82,4.44,10.20,.64,'在 Fe–30.34Mn–8.93Al–0.9C 钢经热轧水淬、未再固溶的状态下，\n为什么残留 κ-碳化物会限制连续致密富 Al 保护膜的形成？',17,NAVY,True)
text(s,1.10,5.38,11.08,.57,'答案组织：碳化物占用 Al → 表面成膜供应受限 → 膜的保护能力下降。\n路径决定问答主题；事实卡提供解释内容，下一页展示实际保留的答案与评分。',16,INK)
banner(s,'实现要点：生成器接收文本化的题型、焦点和事实卡，不直接把整张图作为输入。',y=6.11,size=14.5)
note('spec','代码怎样把子图变成问答','代码先遍历机制路径，并根据路径两端节点找到所属样品。对 path-5，出题规格的 qa_type 是 mechanism_forward，anchor_ids 是 r5 和 r9，relation_path 保存两段节点关系，focus 使用路径的 mechanism 文本。接着构造事实卡，补充成分、工艺、Wiki 机制以及研究对象说明，并合并路径证据、样品证据和解释性段落引用。generate_variants 实际发送给模型的是题型、出题焦点、事实卡和要求生成的条数；不是把整张图或完整 QaSpec JSON 直接发过去。机制解释题要求围绕原因和后果提问，于是这条路径被转成为什么 κ-碳化物限制保护膜形成的问题。S0 被完整的成分与工艺描述替代，r5 等内部代号也不会进入训练问答。答案再根据事实卡解释 Al 被占用、成膜供应受限和保护性下降。页面问题为保留样例的精简展示，完整原文在配套文件中。',['sft_mechanism_qa_pipline/sft_mechanism_qa/qa/spec.py：_path_sample_and_labels / build_specs','sft_mechanism_qa_pipline/sft_mechanism_qa/qa/generate.py：generate_variants','source/qa_case.json'])

# Platform references plus concrete data cards.
s=page('platforms','平台借鉴：材料条目、分层与检索','Materials Project 的条目组织已作为参考；企业与 AI 平台经验用于进一步设计对照。','来源：MP 计算数据快照（Petousis et al., 2017）；Databricks / Pinecone 官方文档；DOI 10.1002/srin.202400475')
rect(s,1.05,2.08,5.43,2.18,LIGHT);text(s,1.23,2.24,5.04,.37,'Materials Project｜SiO₂ · mp-7029',19,BLUE,True)
text(s,1.23,2.80,5.02,.94,'结构：空间群 96；计算胞含 12 个原子\nPBE 带隙：5.51 eV；总介电常数：4.04\n性质与结构归属于同一 material_id',15,INK)
text(s,1.23,3.85,5.02,.22,'2017 年公开计算数据快照（matminer 收录）；非实验测量值。',10.5,MUTED)
rect(s,6.79,2.08,5.43,2.18,PALETEAL);text(s,6.97,2.24,5.05,.37,'SteelDig｜低密度钢 · 铸态样品',21,TEAL,True)
text(s,6.97,2.80,5.02,.31,'Fe–12Mn–9Al–3Cr–1.4C–0.02Nb',16,INK,True)
text(s,6.97,3.23,5.02,.55,'屈服：880 ± 19 MPa    抗拉：1016 ± 12 MPa\n延伸率：39 ± 2 %    证据：b0035 / b0049',14,INK)
text(s,6.97,3.94,5.02,.20,'本地真实抽取记录；同成分不同工艺状态分别建样品。',10.5,MUTED)
arrow(s,6.51,3.12,6.75,3.12)
for x,t,b,c in [(1.05,'MP｜已参考','稳定 ID、元素筛选、详情与 API\n本库以“论文＋样品”替代结构 ID。',BLUE),(4.88,'Databricks｜设计对照','原始 → 清洗 → 服务的分层思想\n借鉴可重建流程；当前未采用该平台。',TEAL),(8.71,'Pinecone｜设计对照','向量记录带元数据，支持条件过滤\n本库使用 pgvector，并保留来源标识。',GOLD)]:
 rect(s,x,4.63,3.56,1.19,LIGHT);text(s,x+.14,4.78,3.28,.34,t,18,c,True);text(s,x+.14,5.22,3.28,.49,b,12.5,INK)
banner(s,'借鉴的是数据组织与治理方法：保留“计算结构”与“文献实验样品”的身份差异。',size=14.5)
note('platforms','具体借鉴哪些平台，数据长什么样','Materials Project 是已有设计文档明确参考的平台。它使用稳定 material_id 组织晶体结构与性质，支持元素筛选、详情页和 API。左侧展示来自 Materials Project 的真实计算数据快照，由 matminer 的 dielectric_constant 数据集收录，对应 Petousis 等人在 2017 年发表的公开数据。条目 mp-7029 为 SiO2，空间群编号 96，计算胞含 12 个原子，PBE 带隙为 5.51 eV，总介电常数的平均值为 4.04。这是注明来源和年代的计算数据，不是当前 API 实时结果，也不是实验测量；其完整 structure、CIF 和性质保存在配套快照。官方 API 文档则用于核对稳定 ID、字段选择与计算来源追踪方式。右侧是本库真实论文样品，Fe–12Mn–9Al–3Cr–1.4C–0.02Nb 铸态钢，屈服强度 880±19 MPa、抗拉强度 1016±12 MPa、延伸率 39±2%，对应文献 DOI 10.1002/srin.202400475。两类条目的身份不同：MP 按晶体结构组织，我们按论文与实验样品组织，相同成分不同工艺也分开。企业平台方面参考 Databricks 的 Medallion 分层理念，但其三层数据精炼流程不能与本库四层一一等同；AI 检索平台方面参考 Pinecone 的向量与元数据设计，本库实际使用的是 pgvector。这两项是本次新增的设计对照，不声称项目历史上采用过这些产品。',['https://doi.org/10.1038/sdata.2016.134','https://ndownloader.figshare.com/files/13213475','https://docs.materialsproject.org/downloading-data/using-the-api/examples','https://docs.materialsproject.org/downloading-data/using-the-api/querying-data','https://docs.databricks.com/aws/en/lakehouse/medallion','https://docs.pinecone.io/guides/search/filter-by-metadata','02_extraction/10.1002_srin.202400475/output/sample/sample_fe-12mn-9al-3cr-1.4c-0.02nb__as-cast.data.json'])

# Explain connections explicitly: build flow on left, keys on right.
s=page('links','四层怎样连通：处理流程与共同标识','四层是逻辑分工：事实表与检索索引共用 PostgreSQL，文件字节存放在本地资产湖。','实现：lake.py / ingest.py / searchidx.py；020_lake.sql / 050_search.sql；API main.py / web.py')
layerrows=[(2.10,'01 资产湖','JSON / 图片 → SHA-256 对象','lake.object_ref：paper_id＋rel_path → object_id\nlake.object：object_id → uri（物理文件）',BLUE),
 (3.07,'02 事实源','ETL 解析、归一、关联后写表','core.*：paper_id＋sample_id 关联样品事实\ntrace：再加 path，连接字段与证据块',TEAL),
 (4.04,'03 检索层','从事实生成检索文本与索引','search.document：doc_id、paper_id、sample_id、ref_id\n文本分词＋全文/模糊索引；向量字段另行填充',GOLD),
 (5.01,'04 消费层','页面 / API 查询并返回结果','精确筛选查事实视图；文本检索查 search.hybrid()\n用共同 ID 打开样品、证据与图表原图',BLUE)]
for y,t,desc,keys,c in layerrows:
 rect(s,1.06,y,4.05,.78,LIGHT);text(s,1.20,y+.10,3.78,.28,t,18,c,True);text(s,1.20,y+.47,3.78,.22,desc,12.5,MUTED)
 arrow(s,5.16,y+.39,5.56,y+.39,c)
 text(s,5.71,y+.11,6.44,.58,keys,13.2,INK)
 if y<5.0:arrow(s,3.08,y+.80,3.08,y+.94,c,1.4)
banner(s,'一次回查：页面 → 样品事实 → path / block_id → 原文证据；图像经 object_id 回到湖。',y=6.10,size=14)
text(s,1.11,6.55,11.1,.18,'事实由 ETL 入库，检索内容由构建函数生成；关联包含数据库外键和应用层标识映射。',10,MUTED)
note('links','四层之间靠什么连接','这页把上一页的概念架构落实到代码。首先，JSON 和图片进入本地资产湖，以 SHA-256 命名。lake.object_ref 把论文号和逻辑路径映射到 object_id，lake.object 再给出物理文件位置。第二，ETL 读取输入，在单篇论文事务中写入 core 表，论文号与样品号把成分、工艺、组织和性能关联起来；字段证据再通过 path 连接到 sample_trace 和 evidence_block。第三，searchidx.build_docs 从论文、Wiki、图表、公式等事实生成检索文档，保留 doc_id、paper_id、sample_id 和 ref_id。中文分词生成检索文本，数据库维护全文和模糊索引，向量字段通过单独的嵌入步骤填充，不是每次修改都天然自动重算。第四，API 用参数化 SQL 查询：数值条件可以直接查事实视图，不必经过向量检索；自然语言问题调用混合检索，结果携带来源标识。图表原图访问则先找到 core.figure.image_paths，再经 lake.object_ref 和 lake.object 取文件。需要区分数据库外键和应用层约定，search.document 的来源字段主要依靠 ETL 与应用保持一致，并非全部声明为外键。',['06_db/etl/steeldig_db/lake.py','06_db/etl/steeldig_db/ingest.py：load_paper','06_db/etl/steeldig_db/searchidx.py：_doc / build_docs','06_db/migrations/020_lake.sql；050_search.sql','06_db/api/steeldig_api/web.py：figure_image / sample_detail'])

# Preserve user's 35 slides, with six inserted in their respective sections.
order=[];origins=[]
for n,s in enumerate(original,1):
 order.append(s);origins.append(n)
 for key in ({2:['background','research'],7:['graph','spec'],16:['platforms'],17:['links']}.get(n,[])):
  order.append(added[key]);origins.append(key)
ids={s.part:e for s,e in zip(list(prs.slides),list(prs.slides._sldIdLst))}
for el in list(prs.slides._sldIdLst):prs.slides._sldIdLst.remove(el)
for s in order:prs.slides._sldIdLst.append(ids[s.part])

# Clarify the existing diagram and evidence example, without altering user's numbers.
for sh in original[6].shapes:
 if sh.has_text_frame and sh.text.startswith('展示链为'):
  sh.text='此页为机制说明；后两页展示真实 path-5 子图及出题过程。'
  for p in sh.text_frame.paragraphs:
   for run in p.runs:run.font.size=Pt(12);run.font.color.rgb=RGBColor.from_string(MUTED)
for sh in original[16].shapes:
 if sh.has_text_frame and sh.text.startswith('分层价值：'):
  sh.text='四层是逻辑分工；事实与检索共用 PostgreSQL。下一页展开处理流程和关联标识。'
  for p in sh.text_frame.paragraphs:
   for run in p.runs:run.font.size=Pt(13);run.font.color.rgb=RGBColor.from_string(NAVY)

# Existing evidence page: show a real field path and identifiers beneath its flow.
s=original[25]
for sh in s.shapes:
 if sh.has_text_frame and sh.text=='样品 + 字段路径\n证据块 b0035 / b0049':
  sh.text='paper_id + sample_id + path\n证据块 b0035 / b0049'
  for p in sh.text_frame.paragraphs:
   for rr in p.runs:rr.font.size=Pt(15)
# Text is editable; this caption occupies the previously empty gap near the bottom.
text(s,1.10,6.02,11.18,.22,'path 实例：$.properties[0].mechanical.tensile_properties.yield_strength[0]  →  block_id',11,MUTED)

# Uniform font, compact title sizing, consistent page numbers; no embedded speaker notes.
for idx,s in enumerate(prs.slides,1):
 for sh in s.shapes:
  if sh.has_text_frame:
   if sh.text.strip().isdigit() and sh.top>Inches(6.6) and sh.left>Inches(11.4):
    p=sh.text_frame.paragraphs[0]
    if p.runs:p.runs[0].text=f'{idx:02d}'
   for p in sh.text_frame.paragraphs:
    p.font.name=FONT
    for rr in p.runs:rr.font.name=FONT
   if idx>1 and sh.top<Inches(.9) and sh.has_text_frame and sh.text.strip():
    sh.text_frame.word_wrap=False
    for p in sh.text_frame.paragraphs:
     for rr in p.runs:
      if rr.font.size and rr.font.size.pt>26:rr.font.size=Pt(26)
  if sh.has_table:
   for row in sh.table.rows:
    for cell in row.cells:
     for p in cell.text_frame.paragraphs:
      p.font.name=FONT
      for rr in p.runs:rr.font.name=FONT
 for rid,rel in list(s.part.rels.items()):
  if rel.reltype.endswith('/notesSlide'):s.part.drop_rel(rid)
prs.save(OUT)
# Set East Asian / Latin fonts in all drawing parts, masters, themes and charts.
with zipfile.ZipFile(OUT) as z:package={n:z.read(n) for n in z.namelist()}
A='http://schemas.openxmlformats.org/drawingml/2006/main'
for name,b in list(package.items()):
 if name.startswith('ppt/') and name.endswith('.xml'):
  try:root=etree.fromstring(b)
  except etree.XMLSyntaxError:continue
  for el in root.iter():
   q=etree.QName(el)
   if q.namespace==A and q.localname in ['latin','ea','cs']:el.set('typeface',FONT)
   if q.namespace==A and q.localname=='font' and el.get('script') in ['Hans','Hant']:el.set('typeface',FONT)
   if q.namespace==A and q.localname in ['rPr','defRPr','endParaRPr']:
    for tag in ['latin','ea','cs']:
     child=el.find('{'+A+'}'+tag)
     if child is None:child=etree.SubElement(el,'{'+A+'}'+tag)
     child.set('typeface',FONT)
  package[name]=etree.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
with zipfile.ZipFile(OUT,'w',zipfile.ZIP_DEFLATED) as z:
 for n,b in package.items():z.writestr(n,b)

# Assemble full standalone manuscript, reusing scripts aligned to retained original slides.
notes=[]
for page_num,origin in enumerate(origins,1):
 if isinstance(origin,str):item=deepcopy(newnotes[origin])
 elif origin==1:item=dict(title='SFT 及数据库架构与应用进展',script='本次汇报分为两部分：第一部分介绍材料领域 SFT 数据建设，从微调背景、真实图谱与问答实例讲到质量审查；第二部分介绍数据库，从平台参考、四层关联讲到查询和治理。模型训练收益尚需独立实验，规模统计仍沿用各页标明的历史快照。',sources=['用户整合后的封面与两部分结构'])
 elif origin==14:item=dict(title='数据库建设',script='下面进入数据库部分。我们把论文解析和抽取资产组织成可查询、可回查、可校订的数据底座。先说明建设结果与平台参考，再解释四层如何通过处理流程和共同标识连接，最后展示查询、证据追踪和维护能力。数据库与 SFT 管线目前各自已有能力，数据库专用的训练导出集成仍属于后续工作。',sources=['README.md；用户整合后的章节页'])
 elif 2<=origin<=13:item=deepcopy(sftnotes[origin-2])
 else:
  base=dbnotes[origin-14];item=dict(title=base['title'],script=base['script'],sources=[base.get('source','')])
 if origin==17:item['script']+=' 四层是逻辑分工，事实表与检索索引当前共用一个 PostgreSQL 实例，下一页用实际字段解释如何连通。'
 if origin==26:item['script']+=' 这条真实样品的标识是 sample_fe-12mn-9al-3cr-1.4c-0.02nb__as-cast，论文号是 10.1002_srin.202400475。字段路径为 $.properties[0].mechanical.tensile_properties.yield_strength[0]，当前代码采用这种 JSONPath 风格路径。property_value 与 sample_trace 在论文号、样品号、run_id 和 path 上对齐，block_id 再在同一论文内定位证据。'
 item['page']=page_num;item['original_page']=origin;notes.append(item)

D=Document();sec=D.sections[0];sec.top_margin=DI(.7);sec.bottom_margin=DI(.7)
for nm in ['Normal','Title','Heading 1','Heading 2']:
 st=D.styles[nm];st.font.name=FONT;st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
 if nm!='Normal':st.font.color.rgb=DC.from_string(NAVY)
normal=D.styles['Normal'];normal.font.size=DP(11);normal.paragraph_format.line_spacing=1.25;normal.paragraph_format.space_after=DP(8)
D.add_heading('SteelDig｜SFT 与数据库汇报',0);D.add_paragraph('优化版独立讲稿 · 对应 41 页 PPT · 2026-10-25')
D.add_paragraph('PPT 内无演讲者注释。新增背景与研究结果附公开来源；现有统计沿用原页历史口径。')
md=['# SteelDig｜SFT 与数据库优化版：独立讲稿','对应 41 页 PPT；PPT 内不写入演讲者注释。']
for n in notes:
 D.add_heading(f"第 {n['page']:02d} 页｜{n['title']}",1);D.add_paragraph(n['script'])
 refs=n.get('sources',[]);refs=[refs] if isinstance(refs,str) else refs
 p=D.add_paragraph('核对依据：'+'；'.join(refs))
 for rr in p.runs:rr.font.size=DP(8.5);rr.font.color.rgb=DC.from_string(MUTED)
 md += [f"\n## 第 {n['page']:02d} 页｜{n['title']}",n['script'],'来源：'+'；'.join(refs)]
D.add_heading('术语速查',1)
for t in ['SFT：用问题与参考答案进行监督训练。','RAG：先检索外部资料，再把证据放进模型上下文生成答案。','持续预训练：在领域文本上继续语言建模训练，与指令 QA 微调不同。','灾难性遗忘：学习新任务后，原有任务能力下降。','QaSpec：生成问答之前确定的题型、焦点、上下文和证据引用。','JSONPath 风格字段路径：在 JSON 中定位某项字段，例如 $.properties[0]。','SHA-256：文件内容的指纹，用于对象命名和校验，不等于备份或权限控制。','RRF：把不同检索方式得到的排名合并。']:
 D.add_paragraph(t)
D.save(ROOT/'SteelDig_SFT与数据库_独立讲稿_1025.docx')
(ROOT/'SteelDig_SFT与数据库_独立讲稿_1025.md').write_text('\n\n'.join(md))
(WORK/'source/slide_notes.json').write_text(json.dumps(notes,ensure_ascii=False,indent=2))
(WORK/'source/slide_mapping.json').write_text(json.dumps([dict(page=i+1,original=v) for i,v in enumerate(origins)],ensure_ascii=False,indent=2))
(WORK/'source/text_boxes.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2))
assert hashlib.sha256(ORIG.read_bytes()).hexdigest()==original_hash
print(OUT);print('Slides:',len(prs.slides));print('Original unchanged:',original_hash)
